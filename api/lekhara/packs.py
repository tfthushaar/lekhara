"""The packs Lekhara prepares for the bank and for K-RERA: the withdrawal pack for the escrow bank,
the quarterly progress report, and the annual audit pack for the CA.

Each pack is built once as headed sections of figures, then written as a PDF to read, print and
sign, or as an Excel workbook to copy figures from. They are prepared from the project's records
for the promoter and the professionals to check; Lekhara files and signs nothing.
"""

import io
from dataclasses import dataclass, field
from datetime import date, timedelta
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Font
from sqlalchemy import select
from sqlalchemy.orm import Session

from lekhara import services, sheets, statutory
from lekhara.engines.escrow import ROLE_CERTIFIES, ROLE_LABEL, ROLE_ORDER, CertifierRole
from lekhara.engines.obligations import _month_end_after
from lekhara.models import (
    DisputeCase,
    EscrowTxn,
    Filing,
    ProgressCertificate,
    Project,
    Receipt,
    Unit,
    WithdrawalRequest,
)
from lekhara.pdf import FIXED, HEADING, NEW_PAGE, write_pdf
from lekhara.rules import KARNATAKA
from lekhara.services import RequestState, ServiceError

ZERO = Decimal("0")
WIDTH = 86


def rupees(value: Decimal | None) -> str:
    """Rs. 1,20,00,000, in Indian grouping, without paise."""
    if value is None:
        return "Not known"
    sign = "-" if value < 0 else ""
    whole = str(int(abs(value).quantize(Decimal("1"))))
    head, tail = whole[:-3], whole[-3:]
    groups = []
    while len(head) > 2:
        groups.insert(0, head[-2:])
        head = head[:-2]
    if head:
        groups.insert(0, head)
    return f"{sign}Rs. {','.join([*groups, tail])}"


def when(d: date | None) -> str:
    return f"{d.day} {d.strftime('%B %Y')}" if d else "Not yet"


@dataclass
class Section:
    title: str
    rows: list[tuple[str, str | Decimal | int]] = field(default_factory=list)
    notes: list[str] = field(default_factory=list)


@dataclass
class Pack:
    title: str
    subtitle: str
    intro: list[str]
    sections: list[Section]
    closing: list[str] = field(default_factory=list)
    extra_pages: list[list[str]] = field(default_factory=list)


def _cell(value) -> str:
    if isinstance(value, Decimal):
        return rupees(value)
    return str(value)


def _row(label: str, value) -> str:
    text = _cell(value)
    space = WIDTH - len(text) - 2
    label = label if len(label) <= space else label[: space - 3] + "..."
    return f"{FIXED}{label} {'.' * (space - len(label))} {text}"


def to_pdf(pack: Pack) -> bytes:
    lines = [f"{HEADING}{pack.title}", pack.subtitle, "", *pack.intro, ""]
    for s in pack.sections:
        lines.append(f"{HEADING}{s.title}")
        lines.extend(_row(label, value) for label, value in s.rows)
        lines.extend(s.notes)
        lines.append("")
    lines.extend(pack.closing)
    for page in pack.extra_pages:
        lines.extend([NEW_PAGE, *page])
    return write_pdf([lines], title=pack.title)


def to_xlsx(pack: Pack) -> bytes:
    book = Workbook()
    sheet = book.active
    sheet.title = "Figures"
    sheets.append(sheet, [pack.title])
    sheet["A1"].font = Font(bold=True, size=13)
    sheets.append(sheet, [pack.subtitle])
    sheet.append([])
    sheet.append(["Section", "Item", "Value"])
    for cell in sheet[4]:
        cell.font = Font(bold=True)
    for s in pack.sections:
        for label, value in s.rows:
            sheets.append(
                sheet,
                [s.title, label, float(value) if isinstance(value, Decimal) else _cell(value)],
            )
            if isinstance(value, Decimal):
                sheet.cell(row=sheet.max_row, column=3).number_format = "#,##,##0"
    sheet.column_dimensions["A"].width = 34
    sheet.column_dimensions["B"].width = 58
    sheet.column_dimensions["C"].width = 22
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


def _promoter(project: Project) -> str:
    org = project.organization
    return org.legal_name or org.name


PREPARED = (
    "Prepared by Lekhara from the project's records for the promoter and the professionals to "
    "check. Lekhara does not file or sign anything; every figure can be traced to its source in "
    "Lekhara."
)


# The withdrawal pack for the escrow bank

STATEMENT = {
    CertifierRole.ARCHITECT: "I have inspected the site of {project} and certify that the physical "
    "progress of the project as on {on} is {pct}.",
    CertifierRole.ENGINEER: "I certify that the cost of development of {project} incurred as on "
    "{on} is {dev}.",
    CertifierRole.CA: "I certify that the total cost incurred and paid on {project} as on {on} is "
    "{total}, and that {draw} may be withdrawn from the project's designated account, in "
    "proportion to the completion of the project.",
}


def withdrawal_pack(state: RequestState, as_of: date) -> Pack:
    r, snap = state.request, state.snapshot
    p, e, money = snap.project, snap.escrow, snap.money
    on = r.submitted_on or as_of
    letter = [
        f"To: The Manager, {p.escrow_bank or 'the bank holding the designated account'}",
        f"From: {_promoter(p)}, promoter of {p.name}",
        f"Project: {p.name}, K-RERA registration {p.krera_registration}",
        f"Date: {when(on)}",
        "",
        f"We request the withdrawal of {rupees(r.amount)} from the designated account of the "
        "project, under section 4(2)(l)(D) of the Real Estate (Regulation and Development) Act, "
        "2016. The amount is in proportion to the completion of the project, as certified by the "
        "project's architect, engineer and chartered accountant, whose certificates follow.",
        "",
        "Enclosed: the certificates of the architect, the engineer and the chartered accountant, "
        "and the position of the designated account with the figures behind the amount.",
        "",
        f"For {_promoter(p)}",
        "",
        "Authorised signatory: ______________________________",
        "Name: ______________________________   Date: ______________",
    ]
    basis = "Cost incurred and paid" if e.cost_basis == "incurred_and_paid" else "Cost incurred"
    sections = [
        Section(
            "How the amount is worked out",
            [
                ("Estimated cost of the project", e.estimated_cost),
                (basis, e.basis_amount),
                ("Proportion of the estimated cost", f"{e.proportion_pct}%"),
                ("Withdrawn to date", e.withdrawn_to_date),
                ("Remaining within the cost incurred", e.limit_remaining),
                ("Balance in the designated account", e.balance),
                ("Can be withdrawn now", e.drawable_now),
                ("Requested in this withdrawal", r.amount),
            ],
        ),
        Section(
            "Deposits into the designated account",
            [
                ("Received from buyers", e.receipts_from_buyers),
                (f"To be deposited ({int(e.deposit_share * 100)}%)", e.required_deposits),
                ("Deposited", e.deposits),
                ("Short", e.deposit_shortfall),
            ],
        ),
        Section(
            "Cost by head",
            [
                ("Land: estimated", money.land.estimated_cost),
                ("Land: incurred", money.land.incurred),
                ("Development: estimated", money.development.estimated_cost),
                ("Development: incurred", money.development.incurred),
            ],
        ),
    ]
    certificate_pages = []
    for role in ROLE_ORDER:
        c = state.certificates.get(role)
        if c is None:
            continue
        figures = {
            "pct": f"{c.completion_pct}%" if c.completion_pct is not None else "not given",
            "dev": rupees(c.development_cost_incurred),
            "total": rupees(c.total_cost_incurred),
            "draw": rupees(c.withdrawable_amount),
        }
        page = [
            f"{HEADING}Certificate of the {ROLE_LABEL[role].lower()}",
            f"{p.name}, K-RERA registration {p.krera_registration}",
            f"Certifies: {ROLE_CERTIFIES[role]}",
            "",
            STATEMENT[role].format(project=p.name, on=when(c.issued_on or on), **figures),
            "",
        ]
        if c.signed:
            page.append(
                f"Signed by {c.signer_name}"
                + (f" of {c.signer_firm}" if c.signer_firm else "")
                + f" on {when(c.issued_on)}."
            )
        else:
            page += [
                f"Not signed yet. Prepared by Lekhara from the books for review and signature by "
                f"the {ROLE_LABEL[role].lower()}, who checks the figures and signs under their "
                "own name and registration.",
                "",
                "Signature: ______________________________",
                "Name and registration number: ______________________________",
                "Date: ______________",
            ]
        certificate_pages.append(page)
    checks = [f"{HEADING}Checks Lekhara ran before the bank sees it"] + [
        f"{c.status.value.upper()}: {c.title}. {c.detail}" for c in state.checks
    ]
    return Pack(
        title="Withdrawal from the designated account",
        subtitle=f"{p.name}, K-RERA registration {p.krera_registration}. As on {when(on)}.",
        intro=letter,
        sections=sections,
        closing=[PREPARED],
        extra_pages=[*certificate_pages, checks],
    )


# The quarterly progress report


@dataclass(frozen=True)
class _Position:
    money: object
    escrow: object


def _position(session: Session, project: Project, on: date) -> _Position:
    """The cost and the designated account as they stood on a date."""
    money = services._money(session, project, on, KARNATAKA)
    return _Position(money, services._escrow(session, project, money, on, KARNATAKA))


def _period_money(session: Session, project: Project, start: date, end: date) -> dict:
    receipts = list(
        session.scalars(
            select(Receipt).where(
                Receipt.project_id == project.id,
                Receipt.received_on <= end,
            )
        )
    )
    txns = list(
        session.scalars(
            select(EscrowTxn).where(EscrowTxn.project_id == project.id, EscrowTxn.txn_date <= end)
        )
    )

    def total(rows, kind=None, *, period: bool, tds: bool = False) -> Decimal:
        out = ZERO
        for t in rows:
            on = t.received_on if isinstance(t, Receipt) else t.txn_date
            if kind is not None and t.kind != kind:
                continue
            if period and on < start:
                continue
            out += t.tds_deducted if tds else t.amount
        return out

    balance = (
        total(txns, "deposit", period=False)
        + total(txns, "interest", period=False)
        - total(txns, "withdrawal", period=False)
        - total(txns, "charges", period=False)
    )
    return {
        "received": total(receipts, period=True),
        "received_to_date": total(receipts, period=False),
        "tds": total(receipts, period=True, tds=True),
        "deposited": total(txns, "deposit", period=True),
        "deposited_to_date": total(txns, "deposit", period=False),
        "withdrawn": total(txns, "withdrawal", period=True),
        "withdrawn_to_date": total(txns, "withdrawal", period=False),
        "interest": total(txns, "interest", period=True),
        "balance": balance,
    }


def _approvals_section(session: Session, project: Project, on: date) -> Section:
    notice = KARNATAKA.value("approvals.renewal_notice_days", on)
    rows = []
    for a in statutory.approvals_of(session, project):
        status = statutory.approval_status(a, as_of=on, notice_days=notice)
        detail = statutory.STATUS_TEXT[status]
        if a.granted_on and a.granted_on <= on:
            detail = f"Granted {a.granted_on.strftime('%d/%m/%Y')}"
            if a.valid_until:
                detail += f", valid to {a.valid_until.strftime('%d/%m/%Y')}"
        elif a.applied_on and a.applied_on <= on:
            detail = f"Applied {a.applied_on.strftime('%d/%m/%Y')}"
        else:
            detail = "Not applied for"
        rows.append((a.name, detail))
    notes = [] if rows else ["No approvals are listed in Lekhara for this project."]
    return Section("Approvals", rows, notes)


def _cases_open(session: Session, project: Project, on: date) -> list[DisputeCase]:
    return [
        c
        for c in session.scalars(select(DisputeCase).where(DisputeCase.project_id == project.id))
        if (c.filed_on or c.notice_received_on or on) <= on
        and (c.closed_on is None or c.closed_on > on)
    ]


def qpr_pack(session: Session, project: Project, quarter_end: date) -> Pack:
    if (quarter_end.month, quarter_end.day) not in {(3, 31), (6, 30), (9, 30), (12, 31)}:
        raise ServiceError("A quarter ends on 31 March, 30 June, 30 September or 31 December.", 404)
    start = date(quarter_end.year, quarter_end.month - 2, 1)
    if quarter_end < project.registered_on:
        raise ServiceError("The project was not registered yet in that quarter.", 404)
    units = list(session.scalars(select(Unit).where(Unit.project_id == project.id)))
    booked = [u for u in units if u.booked and u.booked_on and u.booked_on <= quarter_end]
    in_quarter = [u for u in booked if u.booked_on >= start]
    developer = [u for u in units if u.owner_share != "landowner"]
    progress = session.scalars(
        select(ProgressCertificate)
        .where(
            ProgressCertificate.project_id == project.id, ProgressCertificate.as_of <= quarter_end
        )
        .order_by(ProgressCertificate.as_of.desc())
    ).first()
    snap = _position(session, project, quarter_end)
    money = _period_money(session, project, start, quarter_end)
    cases = _cases_open(session, project, quarter_end)
    filed = session.scalars(
        select(Filing).where(
            Filing.project_id == project.id, Filing.kind == "qpr", Filing.period_end == quarter_end
        )
    ).first()
    due = quarter_end + timedelta(days=KARNATAKA.value("krera.qpr_due_days", quarter_end))
    sections = [
        Section(
            "The project",
            [
                ("Promoter", _promoter(project)),
                ("K-RERA registration", project.krera_registration),
                ("Registered on", project.registered_on.strftime("%d/%m/%Y")),
                (
                    "Possession committed",
                    project.possession_due_on.strftime("%d/%m/%Y")
                    if project.possession_due_on
                    else "Not recorded",
                ),
            ],
        ),
        Section(
            "Units",
            [
                ("Units in the project", len(units)),
                ("The promoter's units", len(developer)),
                ("The landowner's units", len(units) - len(developer)),
                ("Booked by the end of the quarter", len(booked)),
                ("Booked during the quarter", len(in_quarter)),
                ("Not yet booked (the promoter's)", len(developer) - len(booked)),
            ],
        ),
        Section(
            "Construction",
            [
                (
                    "Physical completion, as last certified",
                    f"{progress.completion_pct}% on {progress.as_of.strftime('%d/%m/%Y')}"
                    if progress
                    else "No certificate yet",
                ),
                ("Certified by", progress.certified_by if progress else "Not yet"),
            ],
        ),
        Section(
            "Money from buyers and the designated account",
            [
                ("Received from buyers in the quarter", money["received"]),
                ("Received from buyers to date", money["received_to_date"]),
                ("TDS buyers deducted in the quarter", money["tds"]),
                ("Deposited in the designated account in the quarter", money["deposited"]),
                ("Deposited to date", money["deposited_to_date"]),
                ("Withdrawn in the quarter", money["withdrawn"]),
                ("Withdrawn to date", money["withdrawn_to_date"]),
                ("Interest credited in the quarter", money["interest"]),
                ("Balance at the end of the quarter", money["balance"]),
            ],
            ["Received means what reached the bank; TDS the buyers deducted is shown apart."],
        ),
        Section(
            "Cost",
            [
                ("Estimated cost of the project", snap.money.anticipated_final_cost),
                ("Land: incurred to date", snap.money.land.incurred),
                ("Development: incurred to date", snap.money.development.incurred),
                ("Total incurred to date", snap.money.incurred),
            ],
        ),
        _approvals_section(session, project, quarter_end),
        Section(
            "Cases",
            [(f"{c.case_number} ({c.complainant})", c.status.replace("_", " ")) for c in cases],
            [] if cases else ["No open cases at the end of the quarter."],
        ),
    ]
    state = (
        f"Recorded as filed on {when(filed.filed_on)}."
        if filed
        else f"Due on the K-RERA portal by {when(due)}."
    )
    return Pack(
        title="Quarterly progress report",
        subtitle=(
            f"{project.name}, K-RERA registration {project.krera_registration}. Quarter from "
            f"{when(start)} to {when(quarter_end)}."
        ),
        intro=[
            f"The figures for this quarter's update on the K-RERA portal. {state} Check each "
            "figure and copy it into the portal's form; the architect's, engineer's and CA's "
            "certificates for the quarter go with it.",
        ],
        sections=sections,
        closing=[PREPARED],
    )


# The annual audit pack


def audit_pack(session: Session, project: Project, year_end: date) -> Pack:
    if (year_end.month, year_end.day) != (3, 31):
        raise ServiceError("A financial year ends on 31 March.", 404)
    if year_end < project.registered_on:
        raise ServiceError("The project was not registered yet in that year.", 404)
    start = date(year_end.year - 1, 4, 1)
    money = _period_money(session, project, start, year_end)
    snap = _position(session, project, year_end)
    before = services._money(session, project, start - timedelta(days=1), KARNATAKA)
    share = snap.escrow.deposit_share
    required = (money["received"] * share).quantize(Decimal("0.01"))
    released = list(
        session.scalars(
            select(WithdrawalRequest)
            .where(
                WithdrawalRequest.project_id == project.id,
                WithdrawalRequest.released_on >= start,
                WithdrawalRequest.released_on <= year_end,
            )
            .order_by(WithdrawalRequest.released_on)
        )
    )
    withdrawals = Section(
        "Withdrawals released in the year, and the certificates behind each",
        [],
    )
    for w in released:
        signers = ", ".join(
            f"{ROLE_LABEL[CertifierRole(c.role)].lower()} {c.signer_name}"
            for c in sorted(w.certificates, key=lambda c: ROLE_ORDER.index(CertifierRole(c.role)))
            if c.signed
        )
        withdrawals.rows.append(
            (
                f"{w.released_on.strftime('%d/%m/%Y')}, {w.bank_reference or 'no bank reference'}",
                w.released_amount or w.amount,
            )
        )
        withdrawals.notes.append(
            f"Released {when(w.released_on)}: signed by {signers or 'nobody recorded'}."
        )
    if not released:
        withdrawals.notes.append("No withdrawals were released in the year.")
    filed = session.scalars(
        select(Filing).where(
            Filing.project_id == project.id,
            Filing.kind == "annual_audit",
            Filing.period_end == year_end,
        )
    ).first()
    label = f"{year_end.year - 1}-{str(year_end.year)[2:]}"
    sections = [
        Section(
            "Money from buyers",
            [
                ("Received from buyers in the year", money["received"]),
                ("Received from buyers to the end of the year", money["received_to_date"]),
                ("TDS buyers deducted in the year", money["tds"]),
            ],
            ["Received means what reached the bank; TDS the buyers deducted is shown apart."],
        ),
        Section(
            "The designated account in the year",
            [
                (f"To be deposited ({int(share * 100)}% of the year's receipts)", required),
                ("Deposited", money["deposited"]),
                ("Short", max(required - money["deposited"], ZERO)),
                ("Withdrawn", money["withdrawn"]),
                ("Interest credited", money["interest"]),
                ("Balance at the end of the year", money["balance"]),
            ],
        ),
        withdrawals,
        Section(
            "Cost incurred",
            [
                ("Land: in the year", snap.money.land.incurred - before.land.incurred),
                (
                    "Development: in the year",
                    snap.money.development.incurred - before.development.incurred,
                ),
                ("Total: in the year", snap.money.incurred - before.incurred),
                ("Total: to the end of the year", snap.money.incurred),
                ("Estimated cost of the project", snap.money.anticipated_final_cost),
                ("Withdrawn to the end of the year", snap.escrow.withdrawn_to_date),
                ("Cost incurred and paid to the end of the year", snap.escrow.basis_amount),
            ],
            [
                "Withdrawals to date should not be more than the cost incurred and paid to date, "
                "in proportion to the completion of the project."
            ],
        ),
    ]
    draft = [
        f"{HEADING}Draft for the chartered accountant",
        f"We have examined the accounts of {project.name} (K-RERA registration "
        f"{project.krera_registration}) of {_promoter(project)} for the financial year {label}. In "
        "our opinion, and on the basis of the records shown to us, the amounts collected for the "
        "project were deposited in the designated account as section 4(2)(l)(D) requires, were "
        "used for the project, and were withdrawn in proportion to the completion of the project.",
        "",
        "Prepared by Lekhara for review and signature by the chartered accountant, who examines "
        "the accounts and signs under their own name and membership number. Change the wording to "
        "what the CA finds.",
        "",
        "Signature: ______________________________",
        "Name and membership number: ______________________________",
        "Firm registration number: ______________   UDIN: ______________",
        "Date: ______________",
    ]
    due = _month_end_after(year_end, KARNATAKA.value("krera.annual_audit_months", year_end))
    state = (
        f"Recorded as filed on {when(filed.filed_on)}."
        if filed
        else f"Accounts to be audited by {when(due)}."
    )
    return Pack(
        title=f"Annual audit, {label}",
        subtitle=(
            f"{project.name}, K-RERA registration {project.krera_registration}. Year from "
            f"{when(start)} to {when(year_end)}."
        ),
        intro=[
            f"The figures the CA needs to audit the project's accounts for the year. {state} "
            "Each figure comes from the records in Lekhara; the CA checks them against the "
            "books and the bank statements.",
        ],
        sections=sections,
        closing=[PREPARED],
        extra_pages=[draft],
    )
