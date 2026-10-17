"""Bank statements as banks export them.

Every bank lays its statement out a little differently: HDFC writes Date, Narration,
Chq./Ref.No., Value Dt, Withdrawal Amt., Deposit Amt., Closing Balance; SBI writes Txn Date,
Value Date, Description, Ref No./Cheque No., Debit, Credit, Balance; ICICI writes Transaction
Remarks and Withdrawal Amount (INR ); Axis writes Tran Date, Chq No, Particulars; others give
one Amount column with Dr or Cr beside it. Title rows (the account holder, the account number,
the period) sit above the table and a summary below it. This module finds the table in an
Excel, old Excel, CSV, text or HTML file, whichever bank it came from, and reads every line.
"""

import csv
import hashlib
import io
import re
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from html.parser import HTMLParser
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from lekhara.records import UNITS, ParseError, parse_value
from lekhara.services import ServiceError

HEADER_SEARCH = 60
CENT = Decimal("0.01")

HOW_TO_DOWNLOAD = (
    "In your net banking, open the account's statement, choose the dates, and download it as "
    "Excel (XLS or XLSX) or CSV rather than PDF. Every bank builders use offers this; for a "
    "current account in corporate banking it is under Statements or Reports. If the bank only "
    "emails a PDF, ask the branch for the same statement in Excel."
)

# The bank, from the IFSC printed above the table (its first four letters) or from its name.
IFSC_BANKS = {
    "HDFC": "HDFC Bank",
    "SBIN": "State Bank of India",
    "ICIC": "ICICI Bank",
    "UTIB": "Axis Bank",
    "KKBK": "Kotak Mahindra Bank",
    "CNRB": "Canara Bank",
    "BARB": "Bank of Baroda",
    "UBIN": "Union Bank of India",
    "KARB": "Karnataka Bank",
    "FDRL": "Federal Bank",
    "INDB": "IndusInd Bank",
    "YESB": "Yes Bank",
    "IDFB": "IDFC FIRST Bank",
    "PUNB": "Punjab National Bank",
    "IDIB": "Indian Bank",
    "BKID": "Bank of India",
    "KVBL": "Karur Vysya Bank",
    "SIBL": "South Indian Bank",
    "CIUB": "City Union Bank",
    "RATN": "RBL Bank",
    "AUBL": "AU Small Finance Bank",
}
BANK_NAMES = [
    ("state bank of india", "State Bank of India"),
    ("union bank of india", "Union Bank of India"),
    ("bank of baroda", "Bank of Baroda"),
    ("bank of india", "Bank of India"),
    ("hdfc", "HDFC Bank"),
    ("icici", "ICICI Bank"),
    ("axis bank", "Axis Bank"),
    ("kotak", "Kotak Mahindra Bank"),
    ("canara", "Canara Bank"),
    ("karnataka bank", "Karnataka Bank"),
    ("federal bank", "Federal Bank"),
    ("indusind", "IndusInd Bank"),
    ("yes bank", "Yes Bank"),
    ("idfc", "IDFC FIRST Bank"),
    ("punjab national", "Punjab National Bank"),
    ("indian bank", "Indian Bank"),
    ("karur vysya", "Karur Vysya Bank"),
    ("south indian bank", "South Indian Bank"),
    ("city union", "City Union Bank"),
    ("rbl bank", "RBL Bank"),
    ("axis", "Axis Bank"),
]

_IFSC = re.compile(r"\b([A-Z]{4})0[A-Z0-9]{6}\b")
_ACCOUNT = re.compile(
    r"(?:a/?c|account)\s*(?:no|number|num)?\.?\s*(?:of\s+[a-z ]+?)?[:\-\s]\s*([0-9xX*]{6,20})",
    re.IGNORECASE,
)
_PERIOD = re.compile(
    r"(?:from|period)\s*:?\s*(\d{1,2}[-/. ][A-Za-z0-9]{2,9}[-/. ]\d{2,4})\s*(?:to|-)\s*:?\s*"
    r"(\d{1,2}[-/. ][A-Za-z0-9]{2,9}[-/. ]\d{2,4})",
    re.IGNORECASE,
)
_OPENING = re.compile(
    r"opening\s+balance\s*[:\-]?\s*(?:rs\.?|inr|₹)?\s*([\d,]+\.?\d*)\s*(cr|dr)?", re.IGNORECASE
)
_OPENING_LINE = re.compile(r"opening\s+balance|balance\s+b/?f|brought\s+forward|b/f\b", re.I)
_FOOTER = re.compile(
    r"closing\s+balance|statement\s+summary|total|end\s+of\s+statement|generated\s+on|"
    r"this\s+is\s+a\s+computer|abbreviations|page\s+\d",
    re.IGNORECASE,
)


@dataclass
class Line:
    index: int
    on: date
    narration: str
    reference: str
    debit: Decimal
    credit: Decimal
    balance: Decimal | None
    value_date: date | None = None
    key: str = ""

    @property
    def amount(self) -> Decimal:
        return self.credit if self.credit else self.debit

    @property
    def is_credit(self) -> bool:
        return self.credit > 0


@dataclass
class Statement:
    bank: str | None
    account: str | None
    """The account number as printed, digits (and any masking) only."""
    period_from: date | None
    period_to: date | None
    opening: Decimal | None
    closing: Decimal | None
    lines: list[Line] = field(default_factory=list)
    columns: dict[str, str] = field(default_factory=dict)
    """Each role read, and the column it came from."""
    reversed_order: bool = False
    balance_breaks: list[int] = field(default_factory=list)
    """Lines whose balance does not follow from the line before."""
    opening_line: bool = False

    @property
    def account_tail(self) -> str | None:
        if not self.account:
            return None
        digits = re.sub(r"\D", "", self.account)
        return digits[-4:] if len(digits) >= 4 else None


# Reading the cells


def _cell(value) -> object:
    if value is None:
        return ""
    if isinstance(value, str):
        return value.replace("\xa0", " ").strip()
    return value


def _rows_from_xlsx(data: bytes) -> list[list]:
    try:
        book = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except (InvalidFileException, BadZipFile, KeyError, OSError, ValueError):
        raise ServiceError(f"This Excel file could not be opened. {HOW_TO_DOWNLOAD}") from None
    sheet = max(book.worksheets, key=lambda s: s.max_row or 0)
    rows = [[_cell(v) for v in r] for r in sheet.iter_rows(values_only=True)]
    book.close()
    return rows


def _rows_from_xls(data: bytes) -> list[list]:
    import xlrd

    try:
        book = xlrd.open_workbook(file_contents=data)
    except xlrd.XLRDError:
        raise ServiceError(f"This Excel file could not be opened. {HOW_TO_DOWNLOAD}") from None
    sheet = max(book.sheets(), key=lambda s: s.nrows)
    rows = []
    for r in range(sheet.nrows):
        row = []
        for c in range(sheet.ncols):
            cell = sheet.cell(r, c)
            if cell.ctype == xlrd.XL_CELL_DATE:
                try:
                    row.append(xlrd.xldate.xldate_as_datetime(cell.value, book.datemode))
                except (ValueError, OverflowError):
                    row.append(cell.value)
            else:
                row.append(_cell(cell.value))
        rows.append(row)
    return rows


class _Tables(HTMLParser):
    """Rows of every table in an HTML page: some banks' "Excel" downloads are web pages."""

    def __init__(self) -> None:
        super().__init__()
        self.rows: list[list] = []
        self._row: list | None = None
        self._cell: list[str] | None = None

    def handle_starttag(self, tag, attrs) -> None:
        if tag == "tr":
            self._row = []
        elif tag in ("td", "th") and self._row is not None:
            self._cell = []
        elif tag == "br" and self._cell is not None:
            self._cell.append(" ")

    def handle_endtag(self, tag) -> None:
        if tag in ("td", "th") and self._row is not None and self._cell is not None:
            self._row.append(re.sub(r"\s+", " ", "".join(self._cell)).strip())
            self._cell = None
        elif tag == "tr" and self._row is not None:
            self.rows.append(self._row)
            self._row = None

    def handle_data(self, data) -> None:
        if self._cell is not None:
            self._cell.append(data)


def _text(data: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-16"):
        try:
            text = data.decode(encoding)
            if "\x00" not in text:
                return text
        except UnicodeDecodeError:
            continue
    return data.decode("cp1252", errors="replace")


def _rows_from_text(text: str) -> list[list]:
    sample = text[:20_000]
    delimiter = max((",", "\t", "|", ";"), key=sample.count)
    return [[_cell(v) for v in r] for r in csv.reader(io.StringIO(text), delimiter=delimiter)]


def read_rows(filename: str, data: bytes) -> list[list]:
    if not data:
        raise ServiceError("The file is empty.")
    name = filename.casefold()
    head = data[:2048].lstrip()
    if data[:4] == b"%PDF" or name.endswith(".pdf"):
        raise ServiceError(f"Lekhara reads statements as Excel or CSV, not PDF. {HOW_TO_DOWNLOAD}")
    if data[:2] == b"PK":
        rows = _rows_from_xlsx(data)
    elif data[:8] == b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1":
        rows = _rows_from_xls(data)
    elif head[:1] == b"<" or b"<table" in head.lower() or b"<html" in head.lower():
        parser = _Tables()
        parser.feed(_text(data))
        rows = parser.rows
    elif name.endswith((".csv", ".txt", ".xls", ".tsv")) or b"," in head or b"\t" in head:
        rows = _rows_from_text(_text(data))
    else:
        raise ServiceError(
            f"Lekhara could not read this file as a bank statement. {HOW_TO_DOWNLOAD}"
        )
    return [r for r in rows if any(v not in ("", None) for v in r)]


# Columns


def _norm(header) -> str:
    return re.sub(r"[^a-z0-9]+", " ", str(header).casefold()).strip()


DIRECTION_HEADERS = {
    "dr cr", "cr dr", "drcr", "crdr", "debit credit", "credit debit", "dr cr indicator",
    "txn type", "transaction type", "type", "d c", "c d",
}  # fmt: skip


def column_role(header) -> str | None:
    h = _norm(header)
    if not h:
        return None
    words = h.split()
    if "balance" in words or h.startswith("bal"):
        return "balance"
    if "date" in words or h.endswith(" dt") or h in ("dt", "txn dt", "tran dt", "value dt"):
        return "value_date" if "value" in words else "date"
    if h in DIRECTION_HEADERS:
        return "direction"
    if any(w.startswith(("withdraw", "debit")) for w in words) or h in (
        "dr",
        "dr amount",
        "debits",
    ):
        return "debit"
    if any(w.startswith(("deposit", "credit")) for w in words) or h in ("cr", "cr amount"):
        return "credit"
    if words[0] in ("amount", "amt") or h in ("transaction amount", "txn amount", "tran amount"):
        return "amount"
    if any(w.startswith(("chq", "cheque")) or w in ("utr", "instrument") for w in words):
        return "reference"
    narrative = ("narration", "description", "particular", "remark", "detail", "narrative")
    if any(w.startswith(narrative) for w in words):
        return "narration"
    if any(w in ("ref", "reference") for w in words):
        return "reference"
    if h in ("tran id", "transaction id", "txn id"):
        return "tran_id"
    return None


def _columns(row: list) -> dict[str, int]:
    found: dict[str, int] = {}
    for i, header in enumerate(row):
        role = column_role(header)
        if role and role not in found:
            found[role] = i
    if "date" not in found and "value_date" in found:
        found["date"] = found.pop("value_date")
    if "reference" not in found and "tran_id" in found:
        found["reference"] = found["tran_id"]
    return found


def _score(found: dict[str, int]) -> int:
    has_money = ("debit" in found and "credit" in found) or "amount" in found
    if "date" not in found or not has_money:
        return 0
    return len(found)


def find_table(rows: list[list]) -> tuple[int, dict[str, int]]:
    best, best_found, best_score = -1, {}, 0
    for i, row in enumerate(rows[:HEADER_SEARCH]):
        found = _columns(row)
        score = _score(found)
        if score > best_score:
            best, best_found, best_score = i, found, score
    if best < 0:
        raise ServiceError(
            "Lekhara could not find the table of transactions in this file: a row with the date, "
            "the narration and the amounts (withdrawals and deposits, or debit and credit). "
            + HOW_TO_DOWNLOAD
        )
    return best, best_found


# Values


def parse_date(value) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, int | float) and 20000 < value < 80000:
        from openpyxl.utils.datetime import from_excel

        return from_excel(value).date()
    text = str(value or "").strip()
    if not text or not re.search(r"\d", text):
        return None
    text = re.split(r"\s+\d{1,2}:\d{2}", text)[0].strip()
    try:
        return parse_value(UNITS.field("booked_on"), text)
    except ParseError:
        return None


def parse_amount(value) -> tuple[Decimal | None, str | None]:
    """An amount and any Dr or Cr written with it: 1,23,456.00, 5,000.00 Cr, (250.00), -."""
    if value is None or value == "":
        return None, None
    if isinstance(value, int | float | Decimal):
        return Decimal(str(value)).quantize(CENT), None
    text = str(value).strip()
    side = None
    m = re.search(r"\(?\b(cr|dr)\b\.?\)?\s*$", text, re.IGNORECASE)
    if m:
        side = m.group(1).casefold()
        text = text[: m.start()]
    negative = text.startswith("(") and text.endswith(")") or text.startswith("-")
    digits = re.sub(r"[^\d.]", "", text)
    if not digits or digits == ".":
        return None, side
    try:
        amount = Decimal(digits).quantize(CENT)
    except InvalidOperation:
        return None, side
    return (-amount if negative else amount), side


def _get(row: list, index: int | None):
    if index is None or index >= len(row):
        return ""
    return row[index]


def _direction(value) -> str | None:
    v = _norm(value)
    if v in ("dr", "d", "debit", "w", "withdrawal", "db"):
        return "dr"
    if v in ("cr", "c", "credit", "deposit", "dep"):
        return "cr"
    return None


def _first(pattern: re.Pattern, texts: list[str]) -> re.Match | None:
    for text in texts:
        m = pattern.search(text)
        if m:
            return m
    return None


def _bank(texts: list[str], filename: str) -> str | None:
    for text in texts:
        m = _IFSC.search(text.upper())
        if m and m.group(1) in IFSC_BANKS and re.search(r"\bifs", text, re.I):
            return IFSC_BANKS[m.group(1)]
    joined = " ".join([*texts, filename]).casefold()
    for needle, name in BANK_NAMES:
        if needle in joined:
            return name
    return None


def read_statement(filename: str, data: bytes) -> Statement:
    rows = read_rows(filename, data)
    at, cols = find_table(rows)
    above = [" ".join(str(v) for v in r if v not in ("", None)) for r in rows[:at]]

    lines: list[Line] = []
    opening_from_line = None
    texts_below: list[str] = []
    ended = False
    for row in rows[at + 1 :]:
        on = parse_date(_get(row, cols["date"]))
        narration = str(_get(row, cols.get("narration")) or "").strip()
        reference = str(_get(row, cols.get("reference")) or "").strip()
        if reference.endswith(".0") and reference[:-2].isdigit():
            reference = reference[:-2]
        debit, credit = Decimal("0"), Decimal("0")
        if "debit" in cols and "credit" in cols:
            d, _ = parse_amount(_get(row, cols["debit"]))
            c, _ = parse_amount(_get(row, cols["credit"]))
            debit, credit = abs(d or 0), abs(c or 0)
        else:
            amount, side = parse_amount(_get(row, cols["amount"]))
            side = _direction(_get(row, cols.get("direction"))) or side
            if amount is not None:
                if side is None:
                    side = "dr" if amount < 0 else "cr"
                if side == "dr":
                    debit = abs(amount)
                else:
                    credit = abs(amount)
        balance, balance_side = parse_amount(_get(row, cols.get("balance")))
        if balance is not None and balance_side == "dr":
            balance = -abs(balance)
        row_text = " ".join(str(v) for v in row if v not in ("", None))

        if on is None:
            # Words in the date column after the transactions (a summary, a note) end the
            # table; a row with only narration under a transaction carries its words on.
            if lines and str(_get(row, cols["date"]) or "").strip():
                ended = True
            wraps = bool(re.search(r"[A-Za-z0-9]", narration)) and not _FOOTER.search(row_text)
            if lines and not ended and not debit and not credit and wraps:
                lines[-1].narration = f"{lines[-1].narration} {narration}".strip()
            elif _OPENING_LINE.search(row_text) and balance is not None:
                opening_from_line = balance
            else:
                texts_below.append(row_text)
            continue
        if ended:
            texts_below.append(row_text)
            continue
        if _OPENING_LINE.search(narration) and not debit and not credit:
            opening_from_line = balance
            continue
        if not debit and not credit:
            continue
        lines.append(
            Line(
                index=len(lines),
                on=on,
                narration=re.sub(r"\s+", " ", narration),
                reference=reference,
                debit=debit,
                credit=credit,
                balance=balance,
                value_date=parse_date(_get(row, cols.get("value_date"))),
            )
        )

    if not lines:
        raise ServiceError("The statement has no transactions in it.")
    reversed_order = len(lines) > 1 and lines[0].on > lines[-1].on
    if reversed_order:
        lines.reverse()
        for i, line in enumerate(lines):
            line.index = i

    texts = above + texts_below
    account_match = _first(_ACCOUNT, texts)
    account = re.sub(r"[^0-9xX*]", "", account_match.group(1)) if account_match else None
    period = _first(_PERIOD, above)
    period_from = parse_date(period.group(1)) if period else None
    period_to = parse_date(period.group(2)) if period else None
    opening_match = _first(_OPENING, texts)
    opening = opening_from_line
    if opening is None and opening_match:
        opening, _ = parse_amount(opening_match.group(1))
        if opening is not None and (opening_match.group(2) or "").casefold() == "dr":
            opening = -opening
    first = lines[0]
    if opening is None and first.balance is not None:
        opening = first.balance - first.credit + first.debit

    breaks = []
    previous = opening
    for line in lines:
        if line.balance is None:
            previous = None
            continue
        if previous is not None and abs(
            previous + line.credit - line.debit - line.balance
        ) > Decimal("0.01"):
            breaks.append(line.index)
        previous = line.balance

    statement = Statement(
        bank=_bank(texts, filename),
        account=account,
        period_from=period_from or lines[0].on,
        period_to=period_to or lines[-1].on,
        opening=opening,
        closing=lines[-1].balance,
        lines=lines,
        columns={role: str(rows[at][i]) for role, i in cols.items()},
        reversed_order=reversed_order,
        balance_breaks=breaks,
        opening_line=opening_from_line is not None,
    )
    _keys(statement)
    return statement


def _keys(statement: Statement) -> None:
    """A key for each line that is the same in any download covering it, so the same line in
    two overlapping statements is recognised: the account, date, amounts, balance and words."""
    seen: dict[str, int] = {}
    account = statement.account_tail or "account"
    for line in statement.lines:
        words = re.sub(r"[^a-z0-9]", "", line.narration.casefold())
        base = (
            f"{account}|{line.on.isoformat()}|{line.debit}|{line.credit}|"
            f"{line.balance if line.balance is not None else ''}|{words}"
        )
        n = seen.get(base, 0)
        seen[base] = n + 1
        digest = hashlib.sha256(f"{base}|{n}".encode()).hexdigest()[:20]
        line.key = f"bank:{account}:{digest}"
