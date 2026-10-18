"""Everything a firm has in Lekhara, in one download.

A ZIP file holding an Excel workbook with a sheet for each kind of record, the same records
as JSON for another system to read, every document filed in the vault (every version), and a
note on what is inside. Open to the owner and finance at any time, even when the workspace
is read-only, so a firm can always take its data with it.
"""

import json
import re
import zipfile
from collections.abc import Callable, Iterator
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal
from io import BytesIO
from tempfile import SpooledTemporaryFile

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Font
from sqlalchemy import select
from sqlalchemy.orm import Session

from lekhara.db import Money, Percent
from lekhara.models import (
    BankAccount,
    BankPayer,
    Bill,
    ChangeEntry,
    ChecklistWaiver,
    Commitment,
    CostLineRow,
    DataRoom,
    DataRoomDocument,
    DataRoomVisit,
    DemandRow,
    DisputeCase,
    DisputeHearing,
    DisputeOrder,
    DisputePayment,
    Document,
    EscrowTxn,
    Filing,
    Landowner,
    LandownerPaymentRow,
    LedgerEntry,
    Membership,
    Organization,
    PaymentMilestone,
    ProgressCertificate,
    Project,
    ProjectAccess,
    Receipt,
    TallyLink,
    TrustedSender,
    Unit,
    UnitHandover,
    User,
    Vendor,
    WithdrawalCertificate,
    WithdrawalRequest,
    utcnow,
)
from lekhara.storage import Storage

FORMAT_VERSION = 1
EXCEL_CELL_LIMIT = 32_767
LEFT_OUT = {"organization_id", "token_hash", "storage_key"}
"""Columns that say nothing to the firm (or would give away a secret)."""
NOT_EXPORTED = {
    "organizations": "on the Firm sheet",
    "users": "people are named on each row and listed in the JSON",
    "sessions": "sign-in secrets",
    "one_time_tokens": "sign-in secrets",
    "invitations": "pending invitations carry a secret link",
    "outbound_emails": "copies of emails Lekhara sent",
    "import_batches": "raw rows of uploaded files, already in the records they made",
    "import_mappings": "remembered column choices",
    "tally_imports": "uploaded Tally files, already in the records they made",
    "document_pages": "text read from each document, which is in the files",
    "bank_statements": "uploaded bank statements, already in the records they made",
    "inbound_messages": "emails received, already in the records and documents they brought",
    "inbound_items": "files in those emails, already in the records and documents they made",
    "jobs": "Lekhara's own background work",
    "connectors": "links to the firm's own computers, holding their keys",
    "gst_imports": "GST returns downloaded from the portal, already in the bills they made",
    "statement_passwords": "sealed passwords, never given out",
}
"""Tables left out of the export, and why."""

WORDS = {
    "id": "ID",
    "krera": "K-RERA",
    "rera": "RERA",
    "gstin": "GSTIN",
    "gst": "GST",
    "pan": "PAN",
    "tds": "TDS",
    "oc": "OC",
    "msme": "MSME",
    "utr": "UTR",
    "pct": "%",
    "jda": "JDA",
    "spv": "SPV",
    "sha256": "SHA-256",
    "sqft": "sq ft",
    "udyam": "Udyam",
    "ekhata": "e-Khata",
    "rcm": "RCM",
    "no": "number",
}


@dataclass(frozen=True)
class Scope:
    """The ids that belong to one firm, to pick out its rows from every table."""

    org: str
    projects: list[str]
    units: list[str]
    landowners: list[str]
    cases: list[str]
    requests: list[str]
    rooms: list[str]


@dataclass(frozen=True)
class Sheet:
    title: str
    model: type
    where: Callable[[Scope], object]


SHEETS = [
    Sheet("Projects", Project, lambda s: Project.organization_id == s.org),
    Sheet("Units", Unit, lambda s: Unit.project_id.in_(s.projects)),
    Sheet("Handover", UnitHandover, lambda s: UnitHandover.unit_id.in_(s.units)),
    Sheet("Payment plan", PaymentMilestone, lambda s: PaymentMilestone.project_id.in_(s.projects)),
    Sheet("Demands", DemandRow, lambda s: DemandRow.project_id.in_(s.projects)),
    Sheet("Buyer payments", Receipt, lambda s: Receipt.project_id.in_(s.projects)),
    Sheet(
        "Progress certificates",
        ProgressCertificate,
        lambda s: ProgressCertificate.project_id.in_(s.projects),
    ),
    Sheet("Cost lines", CostLineRow, lambda s: CostLineRow.project_id.in_(s.projects)),
    Sheet("Commitments", Commitment, lambda s: Commitment.project_id.in_(s.projects)),
    Sheet("Cost ledger", LedgerEntry, lambda s: LedgerEntry.project_id.in_(s.projects)),
    Sheet("RERA account", EscrowTxn, lambda s: EscrowTxn.project_id.in_(s.projects)),
    Sheet("Suppliers", Vendor, lambda s: Vendor.organization_id == s.org),
    Sheet("Suppliers' bills", Bill, lambda s: Bill.project_id.in_(s.projects)),
    Sheet("Landowners", Landowner, lambda s: Landowner.project_id.in_(s.projects)),
    Sheet(
        "Landowner payments",
        LandownerPaymentRow,
        lambda s: LandownerPaymentRow.landowner_id.in_(s.landowners),
    ),
    Sheet("Withdrawals", WithdrawalRequest, lambda s: WithdrawalRequest.project_id.in_(s.projects)),
    Sheet(
        "Withdrawal certificates",
        WithdrawalCertificate,
        lambda s: WithdrawalCertificate.request_id.in_(s.requests),
    ),
    Sheet("Filings", Filing, lambda s: Filing.project_id.in_(s.projects)),
    Sheet("Cases", DisputeCase, lambda s: DisputeCase.project_id.in_(s.projects)),
    Sheet("Hearings", DisputeHearing, lambda s: DisputeHearing.case_id.in_(s.cases)),
    Sheet("Orders", DisputeOrder, lambda s: DisputeOrder.case_id.in_(s.cases)),
    Sheet("Payments under orders", DisputePayment, lambda s: DisputePayment.case_id.in_(s.cases)),
    Sheet("Documents", Document, lambda s: Document.organization_id == s.org),
    Sheet(
        "Checklist waivers", ChecklistWaiver, lambda s: ChecklistWaiver.project_id.in_(s.projects)
    ),
    Sheet("Data rooms", DataRoom, lambda s: DataRoom.project_id.in_(s.projects)),
    Sheet(
        "Shared in data rooms", DataRoomDocument, lambda s: DataRoomDocument.room_id.in_(s.rooms)
    ),
    Sheet("Data room visits", DataRoomVisit, lambda s: DataRoomVisit.room_id.in_(s.rooms)),
    Sheet("Tally choices", TallyLink, lambda s: TallyLink.project_id.in_(s.projects)),
    Sheet("Bank accounts", BankAccount, lambda s: BankAccount.project_id.in_(s.projects)),
    Sheet("Bank payers", BankPayer, lambda s: BankPayer.project_id.in_(s.projects)),
    Sheet("Trusted senders", TrustedSender, lambda s: TrustedSender.organization_id == s.org),
    Sheet("Team", Membership, lambda s: Membership.organization_id == s.org),
    Sheet("Certifiers", ProjectAccess, lambda s: ProjectAccess.project_id.in_(s.projects)),
    Sheet("Change history", ChangeEntry, lambda s: ChangeEntry.organization_id == s.org),
]


def heading(column: str) -> str:
    words = [WORDS.get(w, w) for w in column.split("_")]
    text = " ".join(words)
    return text[0].upper() + text[1:]


def _ids(session: Session, column, where) -> list[str]:
    return list(session.scalars(select(column).where(where)))


def scope_for(session: Session, org: Organization) -> Scope:
    projects = _ids(session, Project.id, Project.organization_id == org.id)
    cases = _ids(session, DisputeCase.id, DisputeCase.project_id.in_(projects))
    return Scope(
        org=org.id,
        projects=projects,
        units=_ids(session, Unit.id, Unit.project_id.in_(projects)),
        landowners=_ids(session, Landowner.id, Landowner.project_id.in_(projects)),
        cases=cases,
        requests=_ids(session, WithdrawalRequest.id, WithdrawalRequest.project_id.in_(projects)),
        rooms=_ids(session, DataRoom.id, DataRoom.project_id.in_(projects)),
    )


class Names:
    """Readable names for the ids a row points at, for the workbook."""

    def __init__(self, session: Session, scope: Scope) -> None:
        def lookup(model, label, where) -> dict[str, str]:
            rows = session.execute(select(model.id, *label).where(where)).all()
            return {r[0]: " ".join(str(x) for x in r[1:] if x) for r in rows}

        self.by_column = {
            "project_id": lookup(Project, [Project.name], Project.id.in_(scope.projects)),
            "unit_id": lookup(Unit, [Unit.label], Unit.id.in_(scope.units)),
            "vendor_id": lookup(Vendor, [Vendor.name], Vendor.organization_id == scope.org),
            "cost_line_id": lookup(
                CostLineRow,
                [CostLineRow.code, CostLineRow.name],
                CostLineRow.project_id.in_(scope.projects),
            ),
            "landowner_id": lookup(Landowner, [Landowner.name], Landowner.id.in_(scope.landowners)),
        }
        members = select(Membership.user_id).where(Membership.organization_id == scope.org)
        seats = select(ProjectAccess.user_id).where(ProjectAccess.project_id.in_(scope.projects))
        self.people = {
            u.id: (u.name, u.email)
            for u in session.scalars(select(User).where(User.id.in_(members.union(seats))))
        }

    def of(self, column: str, value: str | None) -> str | None:
        if value is None:
            return None
        if column == "user_id" or column.endswith("_by_user_id") or column.endswith("_user_id"):
            person = self.people.get(value)
            return f"{person[0]} ({person[1]})" if person else value
        names = self.by_column.get(column)
        return names.get(value, value) if names is not None else None


def _columns(model) -> list:
    return [c for c in model.__table__.columns if c.name not in LEFT_OUT]


def _json_value(value):
    if isinstance(value, Decimal):
        return format(value, "f")
    if isinstance(value, datetime | date):
        return value.isoformat()
    return value


def _cell_value(value):
    if isinstance(value, str):
        value = ILLEGAL_CHARACTERS_RE.sub("", value)
        return value[:EXCEL_CELL_LIMIT]
    return value


@dataclass
class Exported:
    filename: str
    file: SpooledTemporaryFile
    counts: dict[str, int]
    documents: int
    missing: list[str]

    def chunks(self, size: int = 1 << 16) -> Iterator[bytes]:
        self.file.seek(0)
        try:
            while chunk := self.file.read(size):
                yield chunk
        finally:
            self.file.close()


def _safe(text: str, limit: int = 80) -> str:
    cleaned = re.sub(r'[\\/:*?"<>|\x00-\x1f]+', " ", text).strip().strip(".")
    return re.sub(r"\s+", " ", cleaned)[:limit] or "Untitled"


def _document_paths(docs: list[Document], projects: dict[str, str]) -> dict[str, str]:
    """A readable, unique path in the ZIP for every version of every document."""
    families: dict[str, int] = {}
    for doc in docs:
        families[doc.family_id] = max(families.get(doc.family_id, 0), doc.version)
    taken: set[str] = set()
    paths = {}
    for doc in sorted(docs, key=lambda d: (d.project_id, d.title, d.version)):
        folder = f"documents/{_safe(projects.get(doc.project_id, 'Firm'))}"
        if doc.archived_on is not None:
            folder += "/archived"
        extension = doc.filename.rsplit(".", 1)[-1].lower() if "." in doc.filename else "bin"
        name = _safe(doc.title)
        if families[doc.family_id] > 1:
            name += f" (version {doc.version})"
        path = f"{folder}/{name}.{extension}"
        n = 2
        while path.lower() in taken:
            path = f"{folder}/{name} ({n}).{extension}"
            n += 1
        taken.add(path.lower())
        paths[doc.id] = path
    return paths


def workspace(
    session: Session, storage: Storage, org: Organization, *, as_of: date, user_id: str | None
) -> Exported:
    scope = scope_for(session, org)
    names = Names(session, scope)
    projects = names.by_column["project_id"]
    records: dict[str, list[dict]] = {}
    book = Workbook(write_only=True)
    bold = Font(bold=True)
    counts: dict[str, int] = {}

    docs = list(session.scalars(select(Document).where(Document.organization_id == org.id)))
    doc_paths = _document_paths(docs, projects)

    firm = book.create_sheet("Firm")
    for label, value in (
        ("Firm", org.name),
        ("Legal name", org.legal_name),
        ("GSTIN", org.gstin),
        ("In Lekhara since", org.created_on),
        ("Exported on", as_of),
    ):
        firm.append(
            [WriteOnlyCell(firm, value=label), WriteOnlyCell(firm, value=_cell_value(value))]
        )
    firm.column_dimensions["A"].width = 20
    firm.column_dimensions["B"].width = 40

    for sheet in SHEETS:
        columns = _columns(sheet.model)
        order = [c for c in columns if c.name == "id"] + [c for c in columns if c.name != "id"]
        rows = list(session.scalars(select(sheet.model).where(sheet.where(scope))))
        counts[sheet.title] = len(rows)
        records[sheet.model.__tablename__] = [
            {c.name: _json_value(getattr(r, c.key)) for c in order} for r in rows
        ]

        ws = book.create_sheet(sheet.title)
        ws.freeze_panes = "A2"
        header = []
        widths = []
        for c in order:
            label = heading(c.name)
            if c.name.endswith("_id") and c.name != "id":
                label = heading(c.name[:-3]) if names.of(c.name, "") is not None else label
            header.append(label)
            widths.append(max(10, min(40, len(label) + 4)))
        if sheet.model is Document:
            header.append("File in this export")
            widths.append(50)
        for i, width in enumerate(widths):
            ws.column_dimensions[_letter(i)].width = width
        ws.append([_bold(ws, h, bold) for h in header])
        for r in rows:
            cells = []
            for c in order:
                value = getattr(r, c.key)
                readable = names.of(c.name, value) if c.name.endswith("_id") else None
                cell = WriteOnlyCell(
                    ws, value=_cell_value(readable if readable is not None else value)
                )
                if isinstance(c.type, Money):
                    cell.number_format = "#,##0.00"
                elif isinstance(c.type, Percent):
                    cell.number_format = "0.00"
                elif isinstance(value, datetime):
                    cell.number_format = "dd mmm yyyy hh:mm"
                elif isinstance(value, date):
                    cell.number_format = "dd mmm yyyy"
                cells.append(cell)
            if sheet.model is Document:
                cells.append(WriteOnlyCell(ws, value=doc_paths.get(r.id)))
            ws.append(cells)

    workbook = BytesIO()
    book.save(workbook)

    # Closed once the response has streamed it (Exported.chunks).
    file = SpooledTemporaryFile(max_size=64 << 20)  # noqa: SIM115
    missing: list[str] = []
    with zipfile.ZipFile(file, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("Lekhara records.xlsx", workbook.getvalue())
        for doc in docs:
            try:
                content = storage.get(doc.storage_key)
            except (FileNotFoundError, ValueError):
                missing.append(doc_paths[doc.id])
                continue
            z.writestr(doc_paths[doc.id], content)
        data = {
            "format": "lekhara-workspace-export",
            "version": FORMAT_VERSION,
            "exported_on": as_of.isoformat(),
            "firm": {
                "id": org.id,
                "name": org.name,
                "legal_name": org.legal_name,
                "gstin": org.gstin,
            },
            "people": [
                {"id": pid, "name": names.people[pid][0], "email": names.people[pid][1]}
                for pid in sorted(names.people)
            ],
            "records": records,
            "files": doc_paths,
            "missing_files": missing,
        }
        z.writestr("Lekhara records.json", json.dumps(data, ensure_ascii=False, indent=1))
        z.writestr("Read me.txt", _readme(org, as_of, counts, len(docs) - len(missing), missing))

    session.add(
        ChangeEntry(
            organization_id=org.id,
            project_id=None,
            record_type="workspace",
            record_id=org.id,
            action="exported",
            changes=json.dumps({"documents": len(docs) - len(missing)}),
            source="typed",
            user_id=user_id,
            at=utcnow(),
        )
    )
    session.commit()
    stamp = as_of.isoformat()
    return Exported(
        filename=f"Lekhara {_safe(org.name, 60)} {stamp}.zip",
        file=file,
        counts=counts,
        documents=len(docs) - len(missing),
        missing=missing,
    )


def _letter(index: int) -> str:
    from openpyxl.utils import get_column_letter

    return get_column_letter(index + 1)


def _bold(ws, value: str, font: Font) -> WriteOnlyCell:
    cell = WriteOnlyCell(ws, value=value)
    cell.font = font
    return cell


def _readme(
    org: Organization, as_of: date, counts: dict[str, int], documents: int, missing: list[str]
) -> str:
    lines = [
        f"Everything {org.name} had in Lekhara on {as_of.strftime('%d %B %Y')}.",
        "",
        "Lekhara records.xlsx  One sheet for each kind of record, with names in place of",
        "                      the ids of projects, units, suppliers and people.",
        "Lekhara records.json  The same records for another system to read, with the ids",
        "                      that tie them together. Amounts are in rupees; dates are",
        "                      year-month-day.",
        "documents/            Every document in the vault, by project, every version.",
        "                      Archived documents are in an archived folder.",
        "",
        "What is inside:",
    ]
    lines += [f"  {title}: {n}" for title, n in counts.items()]
    lines.append(f"  Files: {documents}")
    if missing:
        lines += ["", "These files could not be found in storage and are not in this export:"]
        lines += [f"  {path}" for path in missing]
    return "\r\n".join(lines) + "\r\n"
