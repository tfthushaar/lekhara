"""GSTR-2B and GSTR-2A as the GST portal gives them.

The portal offers GSTR-2B (the monthly statement of input credit, built from what suppliers filed)
as JSON or Excel, and GSTR-2A as JSON. This module reads either into one list of suppliers'
invoices and notes. The JSON is read by its keys wherever they sit: each supplier's GSTIN (ctin)
and trade name (trdnm) with its invoices (inum, dt, val, and txval with igst, cgst, sgst and cess,
or the older item details with iamt, camt, samt and csamt), whether credit is available (itcavl)
and whether reverse charge applies (rev); credit and debit notes (nt, ntnum, ntdt) likewise. The
Excel is read from its B2B and credit note sheets, whose headers run over two rows, by what each
column is called.
"""

import io
import json
import re
from dataclasses import dataclass
from datetime import date
from decimal import Decimal, InvalidOperation
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from lekhara.bank.reader import parse_date
from lekhara.services import ServiceError

ZERO = Decimal("0")
_GSTIN = re.compile(r"\d{2}[A-Z0-9]{13}")
HOW_TO = (
    "On the GST portal, open Returns, then Returns Dashboard, choose the period and GSTR-2B, "
    "View, and download it as JSON or Excel."
)


@dataclass
class Invoice:
    supplier_gstin: str
    supplier_name: str | None
    number: str
    on: date | None
    kind: str
    """invoice, credit_note or debit_note."""
    value: Decimal
    taxable: Decimal
    igst: Decimal
    cgst: Decimal
    sgst: Decimal
    cess: Decimal
    reverse_charge: bool
    itc_available: bool | None
    supplier_period: str | None = None

    @property
    def gst(self) -> Decimal:
        return self.igst + self.cgst + self.sgst + self.cess

    @property
    def key(self) -> str:
        return f"{self.supplier_gstin}|{fold_number(self.number)}|{self.kind}"


@dataclass
class Return:
    form: str
    gstin: str | None
    period: str | None
    invoices: list[Invoice]


def fold_number(number: str) -> str:
    """An invoice number for matching: case, spaces, slashes, dashes and leading zeros aside."""
    text = re.sub(r"[^A-Z0-9]", "", str(number).upper())
    return re.sub(r"(?<![0-9])0+(?=[0-9])", "", text) or text


def _money(value) -> Decimal:
    if value in (None, ""):
        return ZERO
    try:
        return Decimal(str(value).replace(",", "").strip() or "0").quantize(Decimal("0.01"))
    except InvalidOperation:
        return ZERO


def _yes(value) -> bool | None:
    if value in (None, ""):
        return None
    return str(value).strip().casefold() in ("y", "yes", "true", "1")


def _get(d: dict, *keys):
    lowered = {str(k).casefold(): v for k, v in d.items()}
    for key in keys:
        if key in lowered and lowered[key] not in (None, ""):
            return lowered[key]
    return None


# JSON


def _amounts(doc: dict) -> tuple[Decimal, Decimal, Decimal, Decimal, Decimal]:
    taxable = _money(_get(doc, "txval"))
    igst = _money(_get(doc, "igst", "iamt"))
    cgst = _money(_get(doc, "cgst", "camt"))
    sgst = _money(_get(doc, "sgst", "samt"))
    cess = _money(_get(doc, "cess", "csamt"))
    items = _get(doc, "itms", "items")
    if isinstance(items, list) and not taxable:
        for item in items:
            detail = _get(item, "itm_det") if isinstance(item, dict) else None
            detail = detail if isinstance(detail, dict) else item
            if not isinstance(detail, dict):
                continue
            taxable += _money(_get(detail, "txval"))
            igst += _money(_get(detail, "iamt", "igst"))
            cgst += _money(_get(detail, "camt", "cgst"))
            sgst += _money(_get(detail, "samt", "sgst"))
            cess += _money(_get(detail, "csamt", "cess"))
    return taxable, igst, cgst, sgst, cess


def _note_kind(doc: dict) -> str:
    typ = str(_get(doc, "typ", "ntty") or "").casefold()
    return "debit_note" if typ in ("d", "debit", "debit note") else "credit_note"


def _suppliers(node, out: list[Invoice], section: str = "") -> None:
    """Walk the JSON for supplier blocks: anything with a ctin and invoices or notes."""
    if isinstance(node, list):
        for item in node:
            _suppliers(item, out, section)
        return
    if not isinstance(node, dict):
        return
    ctin = _get(node, "ctin")
    docs = _get(node, "inv")
    notes = _get(node, "nt")
    if ctin and (isinstance(docs, list) or isinstance(notes, list)):
        name = _get(node, "trdnm", "lglnm")
        period = _get(node, "supprd", "fp")
        for doc in docs or []:
            if not isinstance(doc, dict):
                continue
            taxable, igst, cgst, sgst, cess = _amounts(doc)
            out.append(
                Invoice(
                    supplier_gstin=str(ctin).strip().upper(),
                    supplier_name=name,
                    number=str(_get(doc, "inum") or "").strip(),
                    on=parse_date(_get(doc, "dt", "idt")),
                    kind="invoice",
                    value=_money(_get(doc, "val")),
                    taxable=taxable,
                    igst=igst,
                    cgst=cgst,
                    sgst=sgst,
                    cess=cess,
                    reverse_charge=bool(_yes(_get(doc, "rev", "rchrg"))),
                    itc_available=_yes(_get(doc, "itcavl")),
                    supplier_period=period,
                )
            )
        for doc in notes or []:
            if not isinstance(doc, dict):
                continue
            taxable, igst, cgst, sgst, cess = _amounts(doc)
            out.append(
                Invoice(
                    supplier_gstin=str(ctin).strip().upper(),
                    supplier_name=name,
                    number=str(_get(doc, "ntnum", "nt_num") or "").strip(),
                    on=parse_date(_get(doc, "dt", "ntdt", "nt_dt")),
                    kind=_note_kind(doc),
                    value=_money(_get(doc, "val")),
                    taxable=taxable,
                    igst=igst,
                    cgst=cgst,
                    sgst=sgst,
                    cess=cess,
                    reverse_charge=bool(_yes(_get(doc, "rev", "rchrg"))),
                    itc_available=_yes(_get(doc, "itcavl")),
                    supplier_period=period,
                )
            )
        return
    for key, value in node.items():
        if isinstance(value, dict | list) and str(key).casefold() not in ("itcsumm", "cpsumm"):
            _suppliers(value, out, str(key))


def read_json(data: bytes) -> Return:
    try:
        payload = json.loads(data.decode("utf-8-sig"))
    except (ValueError, UnicodeDecodeError):
        raise ServiceError(f"This file is not readable JSON. {HOW_TO}") from None
    root = payload.get("data", payload) if isinstance(payload, dict) else payload
    invoices: list[Invoice] = []
    _suppliers(root, invoices)
    if not invoices:
        raise ServiceError(f"There are no suppliers' invoices in this file. {HOW_TO}")
    gstin = _get(root, "gstin") if isinstance(root, dict) else None
    period = _get(root, "rtnprd", "fp", "ret_period") if isinstance(root, dict) else None
    form = "GSTR-2B" if isinstance(root, dict) and _get(root, "docdata", "itcsumm") else "GSTR-2A"
    return Return(form=form, gstin=gstin, period=str(period) if period else None, invoices=invoices)


# Excel

COLUMNS = (
    ("gstin", r"gstin"),
    ("name", r"trade|legal name"),
    ("number", r"(invoice|note) (number|no)"),
    ("kind", r"(note|invoice) type"),
    ("date", r"(invoice|note) date"),
    ("value", r"(invoice|note) value"),
    ("taxable", r"taxable value"),
    ("igst", r"integrated"),
    ("cgst", r"central"),
    ("sgst", r"state"),
    ("cess", r"\bcess\b"),
    ("reverse", r"reverse charge"),
    ("itc", r"itc availab"),
    ("period", r"period"),
)


def _sheet_rows(sheet) -> list[list]:
    return [list(r) for r in sheet.iter_rows(values_only=True)]


def _columns(rows: list[list]) -> tuple[int, dict[str, int]]:
    """Where the table starts, and each column. The portal's headers run over two rows (a
    merged heading above, such as Invoice details or Tax Amount, and the column below)."""
    for i, row in enumerate(rows[:12]):
        if not any("gstin" in str(v or "").casefold() for v in row):
            continue
        below = rows[i + 1] if i + 1 < len(rows) else []
        below_is_data = any(_GSTIN.fullmatch(str(v or "").strip()) for v in below)
        width = max(len(row), 0 if below_is_data else len(below))
        names = []
        for c in range(width):
            top = str(row[c] or "") if c < len(row) else ""
            sub = "" if below_is_data or c >= len(below) else str(below[c] or "")
            names.append(f"{top} {sub}".strip().casefold())
        found: dict[str, int] = {}
        for c, name in enumerate(names):
            for key, pattern in COLUMNS:
                if key not in found and re.search(pattern, name):
                    found[key] = c
                    break
        if "gstin" in found and "number" in found:
            return i + (1 if below_is_data else 2), found
    return -1, {}


def read_excel(data: bytes) -> Return:
    try:
        book = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except (InvalidFileException, BadZipFile, KeyError, OSError, ValueError):
        raise ServiceError(f"This Excel file could not be opened. {HOW_TO}") from None
    invoices: list[Invoice] = []
    gstin = period = None
    for sheet in book.worksheets:
        title = sheet.title.casefold()
        rows = _sheet_rows(sheet)
        if "read me" in title or "summary" in title:
            for row in rows[:20]:
                text = " ".join(str(v) for v in row if v is not None)
                m = re.search(r"\b(\d{2}[A-Z]{5}\d{4}[A-Z][0-9A-Z]Z[0-9A-Z])\b", text)
                if m and not gstin:
                    gstin = m.group(1)
                m = re.search(r"(?:tax period|period)\D*(\w+\s*\d{4}|\d{6})", text, re.I)
                if m and not period:
                    period = m.group(1)
            continue
        if not re.search(r"b2b|cdnr", title) or re.search(r"b2ba|cdnra|amend", title):
            continue
        start, cols = _columns(rows)
        if start < 0:
            continue
        notes = "cdnr" in title
        for row in rows[start:]:

            def cell(key, row=row, cols=cols):
                i = cols.get(key)
                return row[i] if i is not None and i < len(row) else None

            supplier = str(cell("gstin") or "").strip().upper()
            number = str(cell("number") or "").strip()
            if not _GSTIN.fullmatch(supplier) or not number:
                continue
            kind = "invoice"
            if notes:
                kind = (
                    "debit_note" if "debit" in str(cell("kind") or "").casefold() else "credit_note"
                )
            invoices.append(
                Invoice(
                    supplier_gstin=supplier,
                    supplier_name=str(cell("name") or "").strip() or None,
                    number=number,
                    on=parse_date(cell("date")),
                    kind=kind,
                    value=_money(cell("value")),
                    taxable=_money(cell("taxable")),
                    igst=_money(cell("igst")),
                    cgst=_money(cell("cgst")),
                    sgst=_money(cell("sgst")),
                    cess=_money(cell("cess")),
                    reverse_charge=bool(_yes(cell("reverse"))),
                    itc_available=_yes(cell("itc")),
                    supplier_period=str(cell("period") or "") or None,
                )
            )
    book.close()
    if not invoices:
        raise ServiceError(f"Lekhara found no B2B invoices in this file. {HOW_TO}")
    return Return(form="GSTR-2B", gstin=gstin, period=period, invoices=invoices)


def looks_like(filename: str, data: bytes) -> bool:
    """A GSTR-2B or GSTR-2A, judged from the start of the file."""
    if data[:2] == b"PK":
        try:
            book = load_workbook(io.BytesIO(data), read_only=True)
            names = [s.casefold() for s in book.sheetnames]
            book.close()
        except (InvalidFileException, BadZipFile, KeyError, OSError, ValueError):
            return False
        return any("b2b" in n for n in names) and (
            "2b" in filename.casefold() or "read me" in " ".join(names) or "itc" in " ".join(names)
        )
    head = data[:5000].decode("utf-8", "ignore").casefold()
    return '"ctin"' in head or ('"docdata"' in head) or ('"b2b"' in head and '"gstin"' in head)


def read_return(filename: str, data: bytes) -> Return:
    if not data:
        raise ServiceError("The file is empty.")
    if data[:2] == b"PK":
        return read_excel(data)
    return read_json(data)
