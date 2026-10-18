"""GSTR-2B and GSTR-2A files laid out the way the GST portal gives them, for tests. Replace or
add real downloads from design partners beside them."""

import io
import json

from openpyxl import Workbook

CHARS = "0123456789ABCDEFGHIJKLMNOPQRSTUVWXYZ"


def gstin(body: str) -> str:
    total = 0
    for i, char in enumerate(body):
        product = CHARS.index(char) * (1 if i % 2 == 0 else 2)
        total += product // 36 + product % 36
    return body + CHARS[(36 - total % 36) % 36]


FIRM = gstin("29AAKFK1234B1Z")
STEEL = gstin("29ABCDE1234F1Z")
RAVI = gstin("29RAVIE5678K1Z")
PLUMB = gstin("29PLUMB9012Q1Z")

# (supplier GSTIN, trade name, invoice number, date, taxable, igst, cgst, sgst)
INVOICES = [
    (STEEL, "BHARAT TMT STEEL", "BTS-0001", "05-09-2026", 1200000, 0, 108000, 108000),
    (STEEL, "BHARAT TMT STEEL", "BTS/2", "12-09-2026", 300000, 0, 27000, 27000),
    (RAVI, "RAVI ELECTRICALS", "RE-77", "06-09-2026", 250000, 0, 22500, 22500),
    (PLUMB, "NEW PLUMBING CO", "NP-1", "07-09-2026", 50000, 9000, 0, 0),
]
NOTES = [(STEEL, "BHARAT TMT STEEL", "CN-4", "20-09-2026", 10000, 0, 900, 900)]


def gstr2b_json(invoices=INVOICES, notes=NOTES, period="092026", firm=FIRM) -> bytes:
    by_supplier: dict = {}
    for ctin, name, inum, dt, txval, igst, cgst, sgst in invoices:
        block = by_supplier.setdefault(
            ctin,
            {"ctin": ctin, "trdnm": name, "supfildt": "11-10-2026", "supprd": period, "inv": []},
        )
        block["inv"].append(
            {
                "inum": inum, "typ": "R", "dt": dt, "val": txval + igst + cgst + sgst, "pos": "29",
                "rev": "N", "itcavl": "Y", "rsn": "", "diffprcnt": 1, "srctyp": "e-Invoice",
                "txval": txval, "igst": igst, "cgst": cgst, "sgst": sgst, "cess": 0,
            }
        )  # fmt: skip
    cdnr = [
        {
            "ctin": ctin,
            "trdnm": name,
            "nt": [
                {
                    "ntnum": num, "typ": "C", "suptyp": "R", "dt": dt, "val": txval + cgst + sgst,
                    "pos": "29", "rev": "N", "itcavl": "Y", "txval": txval, "igst": igst,
                    "cgst": cgst, "sgst": sgst, "cess": 0,
                }
            ],
        }
        for ctin, name, num, dt, txval, igst, cgst, sgst in notes
    ]  # fmt: skip
    payload = {
        "chksum": "abc",
        "data": {
            "gstin": firm,
            "rtnprd": period,
            "version": "1.0",
            "gendt": "14-10-2026",
            "itcsumm": {"itcavl": {"nonrevsup": {"igst": 0, "cgst": 0, "sgst": 0, "cess": 0}}},
            "docdata": {"b2b": list(by_supplier.values()), "cdnr": cdnr},
        },
    }
    return json.dumps(payload).encode()


def gstr2a_json(invoices=INVOICES, firm=FIRM) -> bytes:
    """The older GSTR-2A layout: amounts in item details."""
    b2b = [
        {
            "ctin": ctin,
            "cfs": "Y",
            "inv": [
                {
                    "inum": inum,
                    "idt": dt,
                    "val": txval + igst + cgst + sgst,
                    "itms": [
                        {
                            "num": 1,
                            "itm_det": {
                                "txval": txval, "iamt": igst, "camt": cgst, "samt": sgst,
                                "csamt": 0, "rt": 18,
                            },
                        }
                    ],
                }
            ],
        }
        for ctin, name, inum, dt, txval, igst, cgst, sgst in invoices
    ]  # fmt: skip
    return json.dumps({"gstin": firm, "fp": "092026", "b2b": b2b}).encode()


def gstr2b_xlsx(invoices=INVOICES, firm=FIRM) -> bytes:
    book = Workbook()
    readme = book.active
    readme.title = "Read me"
    readme.append(["Goods and Services Tax - GSTR 2B"])
    readme.append(["Financial Year", "2026-27"])
    readme.append(["Tax Period", "September 2026"])
    readme.append(["GSTIN", firm])
    sheet = book.create_sheet("B2B")
    sheet.append(["Taxable inward supplies received from registered persons"])
    sheet.append([])
    sheet.append(
        ["GSTIN of supplier", "Trade/Legal name", "Invoice Details", None, None, None,
         "Place of supply", "Supply Attract Reverse Charge", "Taxable Value (₹)", "Tax Amount",
         None, None, None, "GSTR-1/IFF/GSTR-5 Period", "GSTR-1/IFF/GSTR-5 Filing Date",
         "ITC Availability", "Reason", "Applicable % of Tax Rate", "Source", "IRN", "IRN Date"]
    )  # fmt: skip
    sheet.append(
        [None, None, "Invoice number", "Invoice type", "Invoice Date", "Invoice Value(₹)", None,
         None, None, "Integrated Tax(₹)", "Central Tax(₹)", "State/UT Tax(₹)", "Cess(₹)"]
    )  # fmt: skip
    for ctin, name, inum, dt, txval, igst, cgst, sgst in invoices:
        sheet.append(
            [ctin, name, inum, "Regular", dt, txval + igst + cgst + sgst, "Karnataka", "No",
             txval, igst, cgst, sgst, 0, "Sep'26", "11/10/2026", "Yes", "", "", "e-Invoice", "", ""]
        )  # fmt: skip
    book.create_sheet("B2BA").append(["Amendments"])
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()
