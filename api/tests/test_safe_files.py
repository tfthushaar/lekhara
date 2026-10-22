"""Exports keep imported text as text, and uploads that would unpack to gigabytes are refused."""

import csv
import io
import zipfile

import pytest
from openpyxl import Workbook, load_workbook

from lekhara import imports, safe_files, sheets
from lekhara.records import UNITS
from lekhara.services import ServiceError

TRAP = '=HYPERLINK("http://attacker.example","Priya")'


def test_text_that_looks_like_a_formula_stays_text_in_excel() -> None:
    book = Workbook()
    sheet = book.active
    sheets.append(sheet, [TRAP, "Priya", 42])
    out = io.BytesIO()
    book.save(out)
    cell = load_workbook(io.BytesIO(out.getvalue())).active["A1"]
    assert cell.data_type == "s" and cell.value == TRAP

    # The Import centre's own exports and templates go through the same path.
    rows = [[TRAP if f.key == "buyer_name" else "" for f in UNITS.fields]]
    saved = load_workbook(io.BytesIO(imports._xlsx(UNITS, rows, notes=False))).active
    column = [f.key for f in UNITS.fields].index("buyer_name") + 1
    assert saved.cell(row=2, column=column).data_type == "s"


def test_csv_escapes_formulas_and_the_import_undoes_it() -> None:
    assert sheets.csv_value(TRAP) == "'" + TRAP
    assert sheets.csv_value("+91 98450 12345") == "'+91 98450 12345"
    assert sheets.csv_value("@SUM(A1)") == "'@SUM(A1)"
    assert sheets.csv_value("-1,250.50") == "-1,250.50"  # a plain number is left alone
    assert sheets.csv_value("Priya") == "Priya" and sheets.csv_value(7) == 7

    labels = [f.label for f in UNITS.fields]
    row = ["" for _ in labels]
    row[labels.index("Unit")] = "A-101"
    row[labels.index("Buyer")] = TRAP
    data = imports._csv(UNITS, [row])
    written = list(csv.reader(io.StringIO(data.decode("utf-8-sig"))))[1]
    assert written[labels.index("Buyer")] == "'" + TRAP
    back = imports.read_rows("units.csv", data)
    assert back[1][labels.index("Buyer")] == TRAP


def _zip(name: str, size: int) -> bytes:
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr(name, b"\0" * size)
    return out.getvalue()


def test_a_decompression_bomb_is_refused_before_it_is_opened() -> None:
    bomb = _zip("xl/sharedStrings.xml", 30 * 1024 * 1024)  # 30 MB of nothing, in about 30 KB
    assert len(bomb) < 200_000
    with pytest.raises(ServiceError, match="unpacks to far more"):
        safe_files.check_zip(bomb)
    with pytest.raises(ServiceError, match="unpacks to far more"):
        imports.read_rows("units.xlsx", bomb)

    # Ordinary spreadsheets, and files that are not zips at all, pass.
    book = Workbook()
    book.active.append(["Unit", "Buyer"])
    out = io.BytesIO()
    book.save(out)
    safe_files.check_zip(out.getvalue())
    safe_files.check_zip(b"Unit,Buyer\nA-101,Priya\n")
