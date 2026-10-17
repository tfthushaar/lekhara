import io
from datetime import date
from decimal import Decimal as D

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from lekhara.imports import find_header, read_rows
from lekhara.records import UNITS, ParseError, match_columns, parse_value

PASSWORD = "a-long-password"
SG = "sarjapur-greens"


# Parsing


@pytest.mark.parametrize(
    ("raw", "value"),
    [
        ("1,20,00,000", D("12000000.00")),
        ("1.2 Cr", D("12000000.00")),
        ("1.2 crores", D("12000000.00")),
        ("45 lakh", D("4500000.00")),
        ("Rs. 95,00,000/-", D("9500000.00")),
        ("₹ 9500000", D("9500000.00")),
        (9200000.0, D("9200000.00")),
    ],
)
def test_amounts_are_read_the_way_builders_write_them(raw, value) -> None:
    assert parse_value(UNITS.field("agreement_value"), raw) == value


@pytest.mark.parametrize(
    "raw", ["15/03/2025", "15-03-25", "2025-03-15", "15 Mar 2025", "15.3.2025"]
)
def test_dates_are_read_day_first(raw) -> None:
    assert parse_value(UNITS.field("booked_on"), raw) == date(2025, 3, 15)


def test_bad_cells_say_what_is_wrong() -> None:
    with pytest.raises(ParseError, match="not an amount"):
        parse_value(UNITS.field("agreement_value"), "ninety lakh")
    with pytest.raises(ParseError, match="Developer or Landowner"):
        parse_value(UNITS.field("owner_share"), "partner")
    assert parse_value(UNITS.field("owner_share"), "LANDOWNER") == "landowner"
    assert parse_value(UNITS.field("buyer_name"), "  ") is None


def test_columns_are_matched_by_name_and_alias() -> None:
    headers = ["Flat No", "Block", "Carpet Area", "Customer Name", "Sale Value", "Remarks"]
    assert match_columns(UNITS, headers) == {
        0: "label",
        1: "tower",
        2: "carpet_area_sqft",
        3: "buyer_name",
        4: "agreement_value",
    }


def test_title_rows_above_the_header_are_skipped() -> None:
    data = "Kaveri Heights - inventory\nAs on 1 Oct\n\nFlat No,Carpet Area,Customer\nA-1,1100,\n"
    rows = read_rows("units.csv", data.encode())
    assert find_header(UNITS, rows) == 2


# Through the API, on a firm's own project


def firm(app) -> tuple[TestClient, str]:
    client = TestClient(app)
    r = client.post(
        "/auth/signup",
        json={
            "name": "Anil",
            "email": "anil@kaverihomes.in",
            "password": PASSWORD,
            "firm": "Kaveri",
        },
    )
    client.headers["Authorization"] = f"Bearer {r.json()['token']}"
    p = client.post(
        "/projects",
        json={
            "name": "Kaveri Heights",
            "krera_registration": "PRM/K/1",
            "registered_on": "2025-01-01",
        },
    )
    return client, p.json()["id"]


CSV = """Kaveri Heights units
Flat No,Block,Carpet Area,Customer Name,Booking Date,Sale Value,Share
A-101,Tower A,1100,Priya Sharma,15/03/2025,92 lakh,Developer
A-102,Tower A,1150,,,,
A-103,Tower A,,,,,
A-104,Tower A,1200,Ravi K,,1 cr,Landowner
A-101,Tower A,1100,Someone Else,01/01/2025,90 lakh,
"""


def upload(client, project, text: str, name: str = "units.csv"):
    return client.post(
        f"/projects/{project}/imports/units",
        files={"file": (name, text.encode() if isinstance(text, str) else text, "text/csv")},
    )


def test_preview_then_apply(app) -> None:
    client, project = firm(app)
    r = upload(client, project, CSV)
    assert r.status_code == 201, r.text
    batch = client.get(f"/imports/{r.json()['id']}").json()
    assert batch["header_row"] == 2
    assert batch["counts"] == {"add": 2, "update": 0, "unchanged": 0, "error": 3}
    errors = {row["row"]: row["messages"][0] for row in batch["rows"] if row["action"] == "error"}
    assert errors[5] == "The carpet area is missing."
    assert errors[6] == "A sold unit needs its booking date."
    assert errors[7] == "Unit A-101 appears more than once in the file."

    applied = client.post(f"/imports/{batch['id']}/apply").json()
    assert applied["status"] == "applied"
    assert (applied["result"]["added"], applied["result"]["errors"]) == (2, 3)
    units = {u["label"]: u for u in client.get(f"/projects/{project}/units").json()["units"]}
    assert set(units) == {"A-101", "A-102"}
    assert units["A-101"]["booked"] is True

    sold = client.get(f"/units/{units['A-101']['id']}").json()
    assert sold["agreement_value"] == "9200000.00"
    assert sold["history"][0]["source"] == "import"


def test_the_same_file_twice_changes_nothing(app) -> None:
    client, project = firm(app)
    first = upload(client, project, CSV).json()["id"]
    client.post(f"/imports/{first}/apply")
    again = client.get(f"/imports/{upload(client, project, CSV).json()['id']}").json()
    assert again["counts"]["add"] == 0 and again["counts"]["update"] == 0
    assert again["counts"]["unchanged"] == 2


def test_an_update_and_its_undo(app) -> None:
    client, project = firm(app)
    client.post(f"/imports/{upload(client, project, CSV).json()['id']}/apply")
    change = "Flat No,Sale Value,Buyer phone\nA-101,93 lakh,9845012345\n"
    batch = client.get(f"/imports/{upload(client, project, change).json()['id']}").json()
    assert batch["counts"]["update"] == 1
    assert {c["field"] for c in batch["rows"][0]["changes"]} == {"Agreement value", "Buyer phone"}
    applied = client.post(f"/imports/{batch['id']}/apply").json()
    assert applied["can_undo"] is True

    unit_id = next(
        u["id"]
        for u in client.get(f"/projects/{project}/units").json()["units"]
        if u["label"] == "A-101"
    )
    assert client.get(f"/units/{unit_id}").json()["agreement_value"] == "9300000.00"
    undone = client.post(f"/imports/{batch['id']}/undo").json()
    assert undone["status"] == "undone"
    restored = client.get(f"/units/{unit_id}").json()
    assert restored["agreement_value"] == "9200000.00" and restored["buyer_phone"] is None


def test_undo_is_refused_once_records_have_moved_on(app) -> None:
    client, project = firm(app)
    batch = upload(client, project, CSV).json()["id"]
    client.post(f"/imports/{batch}/apply")
    unit_id = next(
        u["id"]
        for u in client.get(f"/projects/{project}/units").json()["units"]
        if u["label"] == "A-102"
    )
    client.put(f"/units/{unit_id}", json={"carpet_area_sqft": "1160"})
    refused = client.post(f"/imports/{batch}/undo")
    assert refused.status_code == 400
    assert "A-102" in refused.json()["detail"]


def test_changing_the_column_mapping(app) -> None:
    client, project = firm(app)
    odd = "Number,Size,Name\nB-1,900,\nB-2,950,\n"
    batch = client.get(f"/imports/{upload(client, project, odd).json()['id']}").json()
    assert batch["missing_required"] == ["Unit", "Carpet area (sq ft)"]
    mapped = client.put(
        f"/imports/{batch['id']}/mapping", json={"mapping": {"0": "label", "1": "carpet_area_sqft"}}
    ).json()
    assert mapped["counts"]["add"] == 2
    twice = client.put(
        f"/imports/{batch['id']}/mapping", json={"mapping": {"0": "label", "1": "label"}}
    )
    assert twice.status_code == 400
    client.post(f"/imports/{batch['id']}/apply")
    # The mapping is remembered for the next file with the same headers.
    third = upload(client, project, "Number,Size,Name\nB-3,1000,\n").json()["id"]
    second = client.get(f"/imports/{third}").json()
    assert second["counts"]["add"] == 1


def test_export_and_template_round_trip(app) -> None:
    client, project = firm(app)
    client.post(f"/imports/{upload(client, project, CSV).json()['id']}/apply")
    exported = client.get(f"/projects/{project}/exports/units.xlsx")
    assert exported.status_code == 200
    sheet = load_workbook(io.BytesIO(exported.content)).active
    assert [c.value for c in sheet[1]][:3] == ["Unit", "Tower", "Carpet area (sq ft)"]
    back = client.post(
        f"/projects/{project}/imports/units",
        files={"file": ("units.xlsx", exported.content, "application/octet-stream")},
    )
    assert client.get(f"/imports/{back.json()['id']}").json()["counts"]["unchanged"] == 2

    template = client.get("/imports/templates/units.xlsx")
    book = load_workbook(io.BytesIO(template.content))
    assert book.sheetnames[1] == "How to fill this in"
    assert (
        client.get("/imports/templates/units.csv").content.decode("utf-8-sig").startswith("Unit,")
    )
    assert client.get("/imports/templates/nothing.xlsx").status_code == 404


def test_excel_files_import(app) -> None:
    client, project = firm(app)
    book = Workbook()
    sheet = book.active
    sheet.append(["Unit", "Carpet area (sq ft)", "Buyer", "Booked on", "Agreement value"])
    sheet.append(["C-1", 1000, "Asha", date(2025, 5, 1), 8500000])
    out = io.BytesIO()
    book.save(out)
    r = client.post(
        f"/projects/{project}/imports/units",
        files={"file": ("units.xlsx", out.getvalue(), "application/octet-stream")},
    )
    batch = client.get(f"/imports/{r.json()['id']}").json()
    assert batch["counts"]["add"] == 1
    junk = upload(client, project, b"not a spreadsheet", name="units.xlsx")
    assert junk.status_code == 400
    assert upload(client, project, "x", name="units.pdf").status_code == 400


def test_typing_a_unit(app) -> None:
    client, project = firm(app)
    r = client.post(
        f"/projects/{project}/units",
        json={"label": "D-1", "carpet_area_sqft": "1,050", "list_price": "1.1 cr"},
    )
    assert r.status_code == 201, r.text
    unit = client.get(f"/units/{r.json()['id']}").json()
    assert unit["list_price"] == "11000000.00"
    assert (
        client.post(
            f"/projects/{project}/units", json={"label": "d-1", "carpet_area_sqft": "900"}
        ).status_code
        == 400
    )
    half_sold = client.put(f"/units/{unit['id']}", json={"buyer_name": "Asha"})
    assert half_sold.status_code == 400
    sold = client.put(
        f"/units/{unit['id']}",
        json={"buyer_name": "Asha", "booked_on": "01/06/2025", "agreement_value": "1.05 cr"},
    ).json()
    assert sold["booked"] is True
    assert len(sold["history"]) == 2
    unsold = client.post(f"/units/{unit['id']}/unsold").json()
    assert unsold["booked"] is False and unsold["agreement_value"] is None
    assert client.post(f"/units/{unit['id']}/delete").json() == {"project_id": project}


def test_units_with_money_against_them_cannot_be_deleted(owner) -> None:
    units = owner.get(f"/projects/{SG}/units").json()["units"]
    sold = next(u for u in units if u["booked"])
    refused = owner.post(f"/units/{sold['id']}/delete")
    assert refused.status_code == 400
    assert "payments received" in refused.json()["detail"]


def test_who_can_import(app, ca) -> None:
    client, _ = firm(app)
    assert upload(client, SG, CSV).status_code == 403  # the sample is read-only
    assert ca.get(f"/projects/{SG}/imports").status_code == 403
