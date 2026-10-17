import json
import zipfile
from datetime import date
from io import BytesIO

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from lekhara.models import ChangeEntry, Organization

PASSWORD = "a-long-password"
PDF = b"%PDF-1.4\n1 0 obj<<>>endobj\ntrailer<<>>\n%%EOF\n"


def firm(app, email="anil@kaverihomes.in") -> tuple[TestClient, str]:
    client = TestClient(app)
    r = client.post(
        "/auth/signup",
        json={"name": "Anil", "email": email, "password": PASSWORD, "firm": "Kaveri Homes"},
    )
    client.headers["Authorization"] = f"Bearer {r.json()['token']}"
    project = client.post(
        "/projects",
        json={
            "name": "Kaveri Heights",
            "krera_registration": "PRM/K/1",
            "registered_on": "2025-04-01",
        },
    ).json()["id"]
    return client, project


def unpack(response) -> tuple[zipfile.ZipFile, dict]:
    assert response.status_code == 200, response.text
    assert response.headers["content-type"] == "application/zip"
    z = zipfile.ZipFile(BytesIO(response.content))
    return z, json.loads(z.read("Lekhara records.json"))


def test_everything_comes_out_in_one_file(app) -> None:
    client, project = firm(app)
    unit = client.post(
        f"/projects/{project}/units",
        json={
            "label": "A-101",
            "carpet_area_sqft": "1100",
            "buyer_name": "Priya Sharma",
            "booked_on": "01/04/2025",
            "agreement_value": "92 lakh",
        },
    )
    assert unit.status_code == 201, unit.text
    client.post(
        f"/units/{unit.json()['id']}/receipts",
        json={"received_on": "05/06/2025", "amount": "4.95 lakh", "reference": "UTR1"},
    )
    replaces: dict = {}
    for version in (1, 2):
        r = client.post(
            f"/projects/{project}/documents",
            data={"title": "Fire NOC", "category": "approval"} | replaces,
            files={"file": ("noc.pdf", PDF + bytes([version]), "application/pdf")},
        )
        assert r.status_code == 201, r.text
        replaces = {"replaces": r.json()["id"]}

    response = client.get("/organization/export")
    assert "Kaveri" in response.headers["content-disposition"]
    z, data = unpack(response)
    names = set(z.namelist())
    assert {"Lekhara records.xlsx", "Lekhara records.json", "Read me.txt"} <= names
    assert "documents/Kaveri Heights/Fire NOC (version 1).pdf" in names
    assert "documents/Kaveri Heights/Fire NOC (version 2).pdf" in names

    records = data["records"]
    assert data["firm"]["name"] == "Kaveri Homes"
    assert [p["name"] for p in records["projects"]] == ["Kaveri Heights"]
    assert records["units"][0]["label"] == "A-101"
    assert records["units"][0]["agreement_value"] == "9200000.00"
    assert records["receipts"][0]["amount"] == "495000.00"
    assert [p["email"] for p in data["people"]] == ["anil@kaverihomes.in"]
    assert all("token_hash" not in row for row in records["data_rooms"])

    book = load_workbook(BytesIO(z.read("Lekhara records.xlsx")))
    assert {"Firm", "Projects", "Units", "Buyer payments", "Documents"} <= set(book.sheetnames)
    receipts = list(book["Buyer payments"].iter_rows(values_only=True))
    header, row = receipts[0], receipts[1]
    assert row[header.index("Unit")] == "A-101"
    assert row[header.index("Project")] == "Kaveri Heights"
    assert float(row[header.index("Amount")]) == 495000
    docs = list(book["Documents"].iter_rows(values_only=True))
    assert docs[0][-1] == "File in this export"

    readme = z.read("Read me.txt").decode()
    assert "Units: 1" in readme and "Files: 2" in readme


def test_only_the_firms_own_data(app) -> None:
    client, _ = firm(app)
    _, data = unpack(client.get("/organization/export"))
    text = json.dumps(data)
    assert "Sarjapur" not in text and "demo.lekhara" not in text
    assert data["records"]["vendors"] == []


def test_who_can_export(app, owner, cfo, accounts, architect) -> None:
    assert owner.get("/organization/export").status_code == 200
    assert cfo.get("/organization/export").status_code == 200
    assert accounts.get("/organization/export").status_code == 403
    assert architect.get("/organization/export").status_code in (403, 404)


def test_a_read_only_workspace_can_still_export(app, engine) -> None:
    from lekhara.db import make_sessionmaker

    client, project = firm(app)
    with make_sessionmaker(engine)() as session:
        org = session.query(Organization).filter_by(name="Kaveri Homes").one()
        org.trial_ends_on = date(2026, 10, 1)
        session.commit()
    assert client.post(f"/projects/{project}/cost-lines/standard").status_code == 402
    z, data = unpack(client.get("/organization/export"))
    assert data["records"]["projects"][0]["name"] == "Kaveri Heights"

    with make_sessionmaker(engine)() as session:
        logged = session.query(ChangeEntry).filter_by(action="exported").all()
        assert len(logged) == 1 and logged[0].record_type == "workspace"


def test_the_sample_firm_exports_everything(owner) -> None:
    z, data = unpack(owner.get("/organization/export"))
    assert len(data["records"]["projects"]) == 3
    assert len(data["records"]["units"]) > 100
    files = [n for n in z.namelist() if n.startswith("documents/")]
    assert len(files) == len(data["files"])


def test_a_file_missing_from_storage_is_named_not_fatal(app) -> None:
    client, project = firm(app)
    r = client.post(
        f"/projects/{project}/documents",
        data={"title": "Title deed", "category": "title"},
        files={"file": ("deed.pdf", PDF, "application/pdf")},
    )
    assert r.status_code == 201, r.text
    for path in app.state.storage.root.rglob("*"):
        if path.is_file():
            path.unlink()
    z, data = unpack(client.get("/organization/export"))
    assert data["missing_files"] == ["documents/Kaveri Heights/Title deed.pdf"]
    readme = z.read("Read me.txt").decode()
    assert "could not be found" in readme and "Title deed.pdf" in readme


def test_every_table_is_exported_or_left_out_on_purpose() -> None:
    from lekhara.db import Base
    from lekhara.export import NOT_EXPORTED, SHEETS

    exported = {sheet.model.__tablename__ for sheet in SHEETS}
    every = set(Base.metadata.tables) - {"alembic_version"}
    assert every - exported - set(NOT_EXPORTED) == set()
    assert exported & set(NOT_EXPORTED) == set()
