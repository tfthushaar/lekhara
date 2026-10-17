import io
from datetime import date
from decimal import Decimal

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from lekhara.db import make_sessionmaker
from lekhara.models import Bill, Organization, Vendor

PASSWORD = "a-long-password"
CHARS = "0123456789ABCDEFGHIJKLMNOPQRSTUVWXYZ"


def gstin(body: str) -> str:
    """A GSTIN with the right check character, from its first 14 characters."""
    total = 0
    for i, char in enumerate(body):
        product = CHARS.index(char) * (1 if i % 2 == 0 else 2)
        total += product // 36 + product % 36
    return body + CHARS[(36 - total % 36) % 36]


STEEL = gstin("29ABCDE1234F1Z")
CEMENT = gstin("29PQRSX5678K1Z")


def firm(app, email="anil@kaverihomes.in") -> tuple[TestClient, str]:
    client = TestClient(app)
    r = client.post(
        "/auth/signup",
        json={"name": "Anil", "email": email, "password": PASSWORD, "firm": "Kaveri Homes"},
    )
    client.headers["Authorization"] = f"Bearer {r.json()['token']}"
    project = client.post(
        "/projects",
        json={
            "name": "Kaveri Heights",
            "krera_registration": "PRM/K/1",
            "registered_on": "2025-04-01",
        },
    ).json()["id"]
    client.post(f"/projects/{project}/cost-lines/standard")
    return client, project


MONEY = ("budget", "committed", "incurred", "paid")


def line(client, project, code) -> dict:
    """A cost line's figures, as numbers."""
    found = next(
        x for x in client.get(f"/projects/{project}/money").json()["lines"] if x["code"] == code
    )
    return {k: Decimal(v) if k in MONEY else v for k, v in found.items()}


def test_suppliers(app) -> None:
    client, project = firm(app)
    r = client.post(
        f"/projects/{project}/suppliers",
        json={"name": "Bharat TMT Steel", "gstin": STEEL.lower(), "msme_category": "small"},
    )
    assert r.status_code == 201, r.text
    s = client.get(f"/suppliers/{r.json()['id']}").json()
    assert s["gstin"] == STEEL and s["pan"] == "ABCDE1234F" and s["state"] == "Karnataka"
    assert s["registered"] and s["msme_label"] == "Small"
    assert s["history"][0]["action"] == "created"

    assert (
        client.post(
            f"/projects/{project}/suppliers", json={"name": "bharat tmt  steel"}
        ).status_code
        == 400
    )
    bad = client.post(
        f"/projects/{project}/suppliers", json={"name": "X", "gstin": "29ABCDE1234F1Z0"}
    )
    assert bad.status_code == 400 and "does not add up" in bad.json()["detail"]
    wrong_pan = client.post(
        f"/projects/{project}/suppliers", json={"name": "Y", "gstin": CEMENT, "pan": "ABCDE1234F"}
    )
    assert "does not match" in wrong_pan.json()["detail"]

    sid = s["id"]
    changed = client.put(
        f"/suppliers/{sid}", json={"udyam_number": "udyam-kr-03-0012345", "written_agreement": "no"}
    ).json()
    assert changed["udyam_number"] == "UDYAM-KR-03-0012345" and not changed["written_agreement"]
    cleared = client.put(f"/suppliers/{sid}", json={"clear": ["gstin"]}).json()
    assert cleared["gstin"] is None and not cleared["registered"]
    assert client.post(f"/suppliers/{sid}/delete").json() == {"ok": True}


def test_bills_carry_cost(app) -> None:
    client, project = firm(app)
    r = client.post(
        f"/projects/{project}/bills",
        json={
            "supplier": "Shree Cement Traders",
            "bill_no": "SCT/101",
            "bill_date": "10/04/2026",
            "taxable_value": "10 lakh",
            "gst_amount": "280000",
            "cost_line": "civil",
            "is_cement": "yes",
        },
    )
    assert r.status_code == 201, r.text
    bill = client.get(f"/bills/{r.json()['id']}").json()
    assert bill["supplier"]["name"] == "Shree Cement Traders" and not bill["supplier"]["registered"]
    assert bill["cost_line"]["code"] == "CIVIL" and bill["accepted_on"] == "2026-04-10"
    assert bill["is_cement"] and bill["source"] == "Lekhara"
    # Residential GST gives no input credit, so the GST is cost too.
    civil = line(client, project, "CIVIL")
    assert civil["incurred"] == Decimal("1280000.00") and civil["paid"] == Decimal("0.00")

    paid = client.put(
        f"/bills/{bill['id']}", json={"paid_on": "20/05/2026", "taxable_value": "9 lakh"}
    )
    assert paid.status_code == 200, paid.text
    civil = line(client, project, "CIVIL")
    assert civil["incurred"] == Decimal("1180000.00") and civil["paid"] == Decimal("1180000.00")
    assert {c["field"] for c in paid.json()["history"][0]["changes"]} == {
        "Taxable value",
        "Paid on",
    }

    moved = client.put(f"/bills/{bill['id']}", json={"cost_line": "MEP"}).json()
    assert moved["cost_line"]["code"] == "MEP"
    assert line(client, project, "CIVIL")["incurred"] == Decimal("0.00")
    assert line(client, project, "MEP")["incurred"] == Decimal("1180000.00")
    client.put(f"/bills/{bill['id']}", json={"clear": ["paid_on"]})
    assert line(client, project, "MEP")["paid"] == Decimal("0.00")

    tax = client.get(f"/projects/{project}/tax").json()
    assert Decimal(tax["procurement"]["total"]) >= Decimal("900000")

    dup = client.post(
        f"/projects/{project}/bills",
        json={
            "supplier": "shree cement traders",
            "bill_no": "sct/101",
            "bill_date": "11/04/2026",
            "taxable_value": "1",
        },
    )
    assert dup.status_code == 400 and "already recorded" in dup.json()["detail"]
    future = client.post(
        f"/projects/{project}/bills",
        json={"supplier": "A", "bill_no": "1", "bill_date": "01/01/2027", "taxable_value": "1"},
    )
    assert "future" in future.json()["detail"]
    nowhere = client.post(
        f"/projects/{project}/bills",
        json={
            "supplier": "A",
            "bill_no": "2",
            "bill_date": "01/04/2026",
            "taxable_value": "1",
            "cost_line": "ZZZ",
        },
    )
    assert "no cost line ZZZ" in nowhere.json()["detail"]

    assert client.post(f"/bills/{bill['id']}/delete").json() == {"project_id": project}
    assert line(client, project, "MEP")["incurred"] == Decimal("0.00")


def test_bills_from_tally_stay(app, engine) -> None:
    client, project = firm(app)
    with make_sessionmaker(engine)() as session:
        org = session.query(Organization).filter_by(name="Kaveri Homes").one()
        vendor = Vendor(organization_id=org.id, name="Tally Supplier", registered=False)
        session.add(vendor)
        session.flush()
        bill = Bill(
            project_id=project,
            vendor_id=vendor.id,
            bill_no="T1",
            bill_date=date(2026, 4, 1),
            accepted_on=date(2026, 4, 1),
            taxable_value=Decimal("100"),
            source_ref="tally:guid-1",
        )
        session.add(bill)
        session.commit()
        bill_id, vendor_id = bill.id, vendor.id
    assert client.get(f"/bills/{bill_id}").json()["source"] == "Tally"
    assert "came from Tally" in client.post(f"/bills/{bill_id}/delete").json()["detail"]
    assert "has bills" in client.post(f"/suppliers/{vendor_id}/delete").json()["detail"]


def test_commitments(app) -> None:
    client, project = firm(app)
    r = client.post(
        f"/projects/{project}/commitments",
        json={
            "reference": "WO/CIV/014",
            "cost_line": "CIVIL",
            "amount": "4.5 crore",
            "committed_on": "01/04/2026",
            "supplier": "Shree Constructions",
            "description": "Structure, towers A and B",
        },
    )
    assert r.status_code == 201, r.text
    assert line(client, project, "CIVIL")["committed"] == Decimal("45000000.00")
    view = client.get(f"/projects/{project}/commitments").json()
    c = view["commitments"][0]
    assert c["supplier"] == "Shree Constructions" and view["total"] == "45000000.00"

    varied = client.put(f"/commitments/{c['id']}", json={"amount": "4.75 crore"}).json()
    assert varied["total"] == "47500000.00"
    assert line(client, project, "CIVIL")["committed"] == Decimal("47500000.00")
    again = client.post(
        f"/projects/{project}/commitments",
        json={
            "reference": "wo/civ/014",
            "cost_line": "CIVIL",
            "amount": "1",
            "committed_on": "01/04/2026",
        },
    )
    assert again.status_code == 400
    client.post(f"/commitments/{c['id']}/delete")
    assert line(client, project, "CIVIL")["committed"] == Decimal("0.00")


def upload(client, project, kind, text) -> dict:
    r = client.post(
        f"/projects/{project}/imports/{kind}",
        files={"file": (f"{kind}.csv", text.encode(), "text/csv")},
    )
    assert r.status_code == 201, r.text
    return client.get(f"/imports/{r.json()['id']}").json()


def counts(view) -> dict:
    return view["counts"]


def actions(view) -> list[str]:
    """Each row's action, in the file's order (the preview lists problems first)."""
    return [row["action"] for row in sorted(view["rows"], key=lambda r: r["row"])]


def test_importing_suppliers_bills_and_orders(app) -> None:
    client, project = firm(app)
    suppliers = (
        "Party Name,GSTIN/UIN,MSME Type,Udyam No\n"
        f"Bharat TMT Steel,{STEEL},Small,UDYAM-KR-03-0012345\n"
        "Ravi Electricals,,Micro,\n"
        f"Bad One,{STEEL[:-1]}0,,\n"
    )
    view = upload(client, project, "suppliers", suppliers)
    assert actions(view) == ["add", "add", "error"]
    client.post(f"/imports/{view['id']}/apply")
    again = upload(client, project, "suppliers", suppliers)
    assert counts(again) == {"add": 0, "update": 0, "unchanged": 2, "error": 1}

    bills = (
        "Supplier GSTIN,Invoice No,Invoice Date,Taxable Amount,Tax Amount,Payment Date,Cost Head\n"
        f'{STEEL},BTS/1,05/04/2026,"12,00,000",216000,10/05/2026,CIVIL\n'
        "Ravi Electricals,RE-77,06/04/2026,2.5 lakh,45000,,MEP\n"
        "New Plumbing Co,NP-1,07/04/2026,50000,9000,,Finishes\n"
    )
    view = upload(client, project, "bills", bills)
    assert actions(view) == ["add", "add", "add"]
    new = next(r for r in view["rows"] if r["row"] == 4)
    assert "will be added as a supplier" in new["messages"][0]
    applied = client.post(f"/imports/{view['id']}/apply").json()
    assert applied["status"] == "applied"
    assert line(client, project, "CIVIL")["paid"] == Decimal("1416000.00")
    assert line(client, project, "MEP")["incurred"] == Decimal("295000.00")
    purchases = client.get(f"/projects/{project}/purchases").json()
    assert len(purchases["bills"]) == 3 and purchases["totals"]["unpaid"] == 2
    assert {s["name"] for s in purchases["suppliers"]} >= {"New Plumbing Co"}

    changed = bills.replace("2.5 lakh,45000,,MEP", "2.5 lakh,45000,15/05/2026,MEP")
    view = upload(client, project, "bills", changed)
    assert counts(view) == {"add": 0, "update": 1, "unchanged": 2, "error": 0}
    client.post(f"/imports/{view['id']}/apply")
    assert line(client, project, "MEP")["paid"] == Decimal("295000.00")
    undone = client.post(f"/imports/{view['id']}/undo")
    assert undone.status_code == 200, undone.text
    assert line(client, project, "MEP")["paid"] == Decimal("0.00")

    exported = client.get(f"/projects/{project}/exports/bills.xlsx")
    sheet = load_workbook(io.BytesIO(exported.content)).active
    rows = list(sheet.iter_rows(values_only=True))
    assert rows[0][:3] == ("Supplier", "Bill number", "Bill date") and len(rows) == 4

    orders = (
        "WO No,Cost Head,Contract Value,WO Date,Contractor,Scope\n"
        "WO/CIV/1,CIVIL,4.5 crore,01/04/2026,Shree Constructions,Structure\n"
        "WO/MEP/1,MEP,80 lakh,02/04/2026,Ravi Electricals,Electrical\n"
    )
    view = upload(client, project, "commitments", orders)
    assert actions(view) == ["add", "add"]
    client.post(f"/imports/{view['id']}/apply")
    assert line(client, project, "MEP")["committed"] == Decimal("8000000.00")
    varied = upload(client, project, "commitments", orders.replace("80 lakh", "85 lakh"))
    assert counts(varied) == {"add": 0, "update": 1, "unchanged": 1, "error": 0}


def test_who_can_change_suppliers_and_bills(app, architect, owner) -> None:
    client, project = firm(app)
    sid = client.post(f"/projects/{project}/suppliers", json={"name": "Bharat TMT Steel"}).json()[
        "id"
    ]
    other = TestClient(app)
    r = other.post(
        "/auth/signup",
        json={"name": "B", "email": "b@other.in", "password": PASSWORD, "firm": "Other"},
    )
    other.headers["Authorization"] = f"Bearer {r.json()['token']}"
    assert other.get(f"/suppliers/{sid}").status_code == 404
    assert (
        architect.post(
            f"/projects/{project}/bills",
            json={"supplier": "X", "bill_no": "1", "bill_date": "01/04/2026", "taxable_value": "1"},
        ).status_code
        == 404
    )
    sample = owner.get("/projects/sarjapur-greens/purchases").json()
    assert sample["bills"] and sample["can_edit"]
    assert client.get("/projects/sarjapur-greens/purchases").json()["can_edit"] is False
