"""The withdrawal pack for the bank, the quarterly progress report and the annual audit pack."""

import io

from openpyxl import load_workbook

from lekhara.packs import rupees
from lekhara.pdf import read_pages
from tests.conftest import signed_in


def pdf_text(response) -> str:
    assert response.status_code == 200, response.text
    assert response.headers["content-type"] == "application/pdf"
    return "\n".join(read_pages(response.content))


def test_rupees_in_indian_grouping() -> None:
    from decimal import Decimal

    assert rupees(Decimal("12000000")) == "Rs. 1,20,00,000"
    assert rupees(Decimal("950")) == "Rs. 950"
    assert rupees(Decimal("-123456.60")) == "-Rs. 1,23,457"


def test_withdrawal_pack(owner) -> None:
    escrow = owner.get("/projects/whitefield-residences/escrow").json()
    request = escrow["requests"][0]
    text = pdf_text(owner.get(f"/withdrawals/{request['id']}/pack.pdf"))
    assert "Withdrawal from the designated account" in text
    assert "section 4(2)(l)(D)" in text
    for heading in (
        "How the amount is worked out",
        "Certificate of the architect",
        "Certificate of the engineer",
        "Certificate of the chartered accountant",
        "Checks Lekhara ran",
    ):
        assert heading in text
    assert "Lekhara does not file or sign anything" in text


def test_quarterly_report_pack(owner) -> None:
    base = "/projects/hebbal-heights/packs/qpr/2026-09-30"
    text = pdf_text(owner.get(f"{base}.pdf"))
    assert "Quarterly progress report" in text
    assert "Booked during the quarter" in text
    assert "Received from buyers in the quarter" in text
    assert "Consent to establish" in text

    book = load_workbook(io.BytesIO(owner.get(f"{base}.xlsx").content))
    sheet = book["Figures"]
    items = {row[1]: row[2] for row in sheet.iter_rows(min_row=5, values_only=True)}
    assert isinstance(items["Received from buyers in the quarter"], float)
    assert items["Units in the project"] == "120" or isinstance(items["Units in the project"], str)

    assert owner.get("/projects/hebbal-heights/packs/qpr/2026-08-31.pdf").status_code == 404
    assert owner.get("/projects/hebbal-heights/packs/other/2026-09-30.pdf").status_code == 404
    assert owner.get("/projects/hebbal-heights/packs/qpr/2020-03-31.pdf").status_code == 404


def test_audit_pack(owner, app) -> None:
    text = pdf_text(owner.get("/projects/whitefield-residences/packs/audit/2026-03-31.pdf"))
    assert "Annual audit, 2025-26" in text
    assert "Withdrawals released in the year" in text
    assert "Draft for the chartered accountant" in text
    assert "UDIN" in text
    wrong = owner.get("/projects/whitefield-residences/packs/audit/2026-06-30.pdf")
    assert wrong.status_code == 404

    architect = signed_in(app, "architect@demo.lekhara")
    response = architect.get("/projects/whitefield-residences/packs/audit/2026-03-31.pdf")
    assert response.status_code == 403
